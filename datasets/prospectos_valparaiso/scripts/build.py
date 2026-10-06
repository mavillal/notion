import pandas as pd, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
D=sys.argv[1]; u=pd.read_pickle(D+'/clean.pkl'); log=pd.read_pickle(D+'/log.pkl')
u=u.astype(object).where(u.notna(),None)
F='Arial'; H=Font(name=F,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1A2A4A'); N=Font(name=F,size=10)
wb=Workbook()
def sheet(ws,df,widths=None):
    ws.append(list(df.columns))
    for r in df.itertuples(index=False): ws.append(list(r))
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(vertical='center',wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=N
    ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
    for i,col in enumerate(df.columns,1):
        w=max([len(str(col))]+[len(str(v)) for v in df[col].head(200) if v is not None])
        ws.column_dimensions[get_column_letter(i)].width=min(max(10,w+2),45)
    ws.row_dimensions[1].height=30
    return ws
# Resumen
ws=wb.active; ws.title='Resumen'
ws['A1']='Base de prospectos B2B – Valparaíso / San Antonio / Casablanca / Curacaví'; ws['A1'].font=Font(name=F,bold=True,size=14)
ws['A2']='Fuente: 22 exportaciones de Google Maps (mapsscraper.ai, plan gratuito) del 05-10-2026. Totales calculados con fórmulas sobre la hoja Base_limpia.'; ws['A2'].font=Font(name=F,italic=True,size=9,color='666666')
n=len(u)+1; R=f"Base_limpia!$G$2:$G${n}"; S=f"Base_limpia!$F$2:$F${n}"; SC=f"Base_limpia!$I$2:$I${n}"
r=4
def hdr(cells):
    global r
    for i,t in enumerate(cells): c=ws.cell(r,1+i,t); c.font=H; c.fill=HF
    r+=1
hdr(['Indicador','Valor'])
for k,f in [('Filas brutas (sin pies de página)',296),('Empresas únicas',f'=COUNTA(Base_limpia!$A$2:$A${n})'),
            ('Con teléfono',f'=COUNTIF(Base_limpia!$J$2:$J${n},"+*")'),('Con web propia',f'=COUNTIF(Base_limpia!$N$2:$N${n},"Sí")'),
            ('Score de completitud promedio',f'=ROUND(AVERAGE({SC}),1)')]:
    ws.cell(r,1,k).font=N; c=ws.cell(r,2,f); c.font=Font(name=F,size=10,color='0000FF' if isinstance(f,int) else '000000'); r+=1
ws.cell(5,3,'Dato fijo: 344 filas leídas − 48 filas de pie de página del scraper').font=Font(name=F,size=9,italic=True,color='666666')
r+=1; hdr(['Relevancia','Empresas','% del total'])
top=r
for k in ['Alta','Media','Baja','Descartar']:
    ws.cell(r,1,k).font=N; ws.cell(r,2,f'=COUNTIF({R},A{r})').font=N; c=ws.cell(r,3,f'=B{r}/$B$6'); c.number_format='0.0%'; c.font=N; r+=1
r+=1; hdr(['Segmento','Total','Alta','Media','Score promedio'])
for k in u['Segmento'].value_counts().index:
    ws.cell(r,1,k).font=N
    ws.cell(r,2,f'=COUNTIF({S},A{r})').font=N
    ws.cell(r,3,f'=COUNTIFS({S},A{r},{R},"Alta")').font=N
    ws.cell(r,4,f'=COUNTIFS({S},A{r},{R},"Media")').font=N
    c=ws.cell(r,5,f'=IFERROR(ROUND(AVERAGEIF({S},A{r},{SC}),1),0)'); c.font=N; r+=1
r+=1; hdr(['Comuna','Total','Alta'])
CM=f"Base_limpia!$O$2:$O${n}"
for k in u['Comuna'].value_counts().index:
    ws.cell(r,1,k).font=N; ws.cell(r,2,f'=COUNTIF({CM},A{r})').font=N; ws.cell(r,3,f'=COUNTIFS({CM},A{r},{R},"Alta")').font=N; r+=1
ws.column_dimensions['A'].width=38
for c in 'BCDE': ws.column_dimensions[c].width=14
# Base y Prospectos
b=sheet(wb.create_sheet('Base_limpia'),u)
fills={'Alta':'D9EAD3','Media':'FFF2CC','Baja':'F4CCCC','Descartar':'D9D9D9'}
for k,v in fills.items():
    b.conditional_formatting.add(f'G2:G{n}',CellIsRule(operator='equal',formula=[f'"{k}"'],fill=PatternFill('solid',fgColor=v)))
for row in b.iter_rows(min_row=2,min_col=25,max_col=25):
    for c in row: c.number_format='0.0'
p=u[u['Relevancia'].isin(['Alta','Media'])]
sheet(wb.create_sheet('Prospectos'),p[['Empresa','Grupo_empresa','Segmento','Relevancia','Score_completitud','Telefono','Tipo_telefono','Sitio_web','Comuna','Direccion','Categoria_principal','N_resenas','Rating','Google_Maps_URL','Flags_calidad']])
sheet(wb.create_sheet('Revisar_descartados'),u[u['Relevancia'].isin(['Baja','Descartar'])][['Empresa','Nombre_original','Relevancia','Motivo_relevancia','Categorias','Direccion','Telefono','Google_Maps_URL']])
dic=pd.DataFrame([
 ('Place_Id','ID único de Google Maps; clave de deduplicación'),('Empresa','Nombre limpio (sin relleno SEO)'),('Nombre_original','Nombre tal cual en Google Maps'),
 ('Grupo_empresa','Empresa matriz cuando varias fichas comparten dominio web o teléfono'),('N_sedes_grupo','Nº de fichas del mismo grupo'),
 ('Segmento','Rubro asignado por reglas sobre categorías de Google, nombre y búsqueda de origen'),
 ('Relevancia','Alta: industrial/logística en zona objetivo con contacto · Media: falta contacto o fuera de zona · Baja: minorista/no industrial · Descartar: no es empresa, duplicado o fuera de Chile'),
 ('Motivo_relevancia','Regla que asignó la relevancia'),
 ('Score_completitud','0–100: teléfono 25 · web propia 25 · dirección con número 15 (aprox. 5) · reseñas ≥50:15, ≥10:10, ≥1:5 · rating ≥4: 10 · horario 10'),
 ('Telefono','Formato internacional +56'),('Tipo_telefono','Móvil / Fijo / 600 / Extranjero / Sin teléfono'),
 ('Web_propia','No si el "sitio" es Instagram, WhatsApp, g.page, Wikipedia, negocio.site o ueniweb'),
 ('Comuna / Region / Codigo_postal','Extraídos de la dirección completa'),('Precision_direccion','Con número / Sin número / Camino sin nombre / Plus Code / Solo comuna / Sin dirección'),
 ('Busquedas_origen','Búsquedas de Google Maps en que apareció la empresa'),('N_apariciones','Veces que la empresa apareció en las 22 exportaciones'),
 ('Flags_calidad','Problemas pendientes de cada ficha'),
 ('Columnas eliminadas','Email y Social Medias (bloqueados por plan gratuito), Claimed (siempre false), Price (1 valor), Phones, Review URL, Featured Image, Cid'),
],columns=['Campo','Descripción'])
d=sheet(wb.create_sheet('Diccionario'),dic); d.column_dimensions['B'].width=110
l=sheet(wb.create_sheet('Log_limpieza'),log); l.column_dimensions['C'].width=110
wb.save(D+'/prospectos_valparaiso_limpio.xlsx')
u.to_csv(D+'/prospectos_valparaiso_limpio.csv',index=False,encoding='utf-8-sig')
