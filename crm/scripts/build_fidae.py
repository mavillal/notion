import csv, re, sys, json, collections, unicodedata
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
# usage: build_fidae.py fidae.csv CRM_Directorio_SSO_2026.xlsx CRM_ExpoMina_2026.xlsx CRM_Exponor_2026.xlsx CRM_AquaSur_2026.xlsx out.xlsx
src, dirx, expox, exponx, aquax, out = sys.argv[1:7]
def clean(s): return re.sub(r'\s+',' ',(s or '').replace('\xa0',' ')).strip()
def norm(s): return unicodedata.normalize('NFKD',s.upper()).encode('ascii','ignore').decode()
def key(n):
    n=norm(n).lower()
    n=re.sub(r'\(.*?\)','',n)
    n=re.sub(r'\b(s\.?\s?a\.?\s?c\.?|s\.?\s?a\.?|s\.?\s?a\.?\s?u\.?|e\.?i\.?r\.?l\.?|s\.?r\.?l\.?|s\.?r\.?o\.?|s\.?l\.?|ltda\.?|limitada|spa|sac|eirl|srl|peru|chile|del|de|y|cia|co|ltd|inc|corp|corporation|company|group|grupo|gmbh|ag|as|llc|plc)\b','',n)
    return re.sub(r'[^a-z0-9]','',n)
def load(path):
    ws=openpyxl.load_workbook(path)['Contactos']; h=[c.value for c in ws[1]]; m={}
    for r in ws.iter_rows(min_row=2,values_only=True):
        d=dict(zip(h,r))
        if d.get('Empresa'): m.setdefault(key(d['Empresa']),d)
    return m
CRMS=[('Directorio SSO',load(dirx)),('ExpoMina',load(expox)),('Exponor',load(exponx)),('AquaSur',load(aquax))]

PAIS={'United States of America':'Estados Unidos','South Korea':'Corea del Sur','Czech Republic':'República Checa','Brazil':'Brasil','Spain':'España',
 'France':'Francia','Austria':'Austria','Germany':'Alemania','Poland':'Polonia','Türkiye':'Turquía','United Kingdom':'Reino Unido','Canada':'Canadá',
 'Switzerland':'Suiza','Italy':'Italia','Belgium':'Bélgica','United Arab Emirates':'Emiratos Árabes Unidos','Greece':'Grecia','Mexico':'México',
 'Finland':'Finlandia','South Africa':'Sudáfrica','Netherlands':'Países Bajos','Singapore':'Singapur','Estonia':'Estonia','Peru':'Perú','Sweden':'Suecia',
 'Ukraine':'Ucrania','Slovakia':'Eslovaquia','Lithuania':'Lituania','Latvia':'Letonia'}
RUB={'Defense (A-010)':'Defensa','Defense':'Defensa','Civil and commercial aviation (A-090)':'Aviación civil y comercial','Others (A-170)':'Otros',
 'Aircraft maintenance (A-110)':'Mantenimiento aeronáutico (MRO)','Space Technologies (A-100)':'Tecnologías espaciales','Homeland Security (A-080)':'Seguridad interior',
 'Aerospace (A-020)':'Aeroespacial','UAV/Drones (A-120)':'UAV / drones','Airport and equipment services (A-050)':'Servicios y equipos aeroportuarios',
 'Cybertech (A-140)':'Ciberseguridad','Media (A-030)':'Medios','Logistics & Ground Support (A-150)':'Logística y apoyo en tierra','Education (A-060)':'Educación','':''}
# Agrupa sub-unidades de instituciones en una sola ficha
GROUP=[(r'^CARABINEROS DE CHILE|^GRUPO DE ADIESTRAMIENTO CANINO','CARABINEROS DE CHILE'),(r'^FUERZA AEREA DE CHILE','FUERZA AÉREA DE CHILE'),
 (r'^POLICIA DE INVESTIGACIONES|^ESUELA DE INVESTIGACIONES','POLICÍA DE INVESTIGACIONES DE CHILE (PDI)'),(r'^SUN ?DRIVE','SUNDRIVE S.R.O.')]
rows=list(csv.DictReader(open(src,encoding='utf-8-sig')))
comp=collections.OrderedDict()
for r in rows:
    raw=clean(r['Empresa']); n=norm(raw); name=raw; unit=''
    for pat,g in GROUP:
        if re.search(pat,n):
            unit=raw if norm(raw)!=norm(g) else ''; name=g; break
    k=key(name)
    c=comp.setdefault(k,{'Empresa':name,'stands':[],'rubros':[],'webs':[],'units':[],'pais':PAIS.get(clean(r['Pais']),clean(r['Pais']))})
    st=clean(r['Ubicacion'])
    if st and st not in c['stands']: c['stands'].append(st)
    rb=RUB.get(clean(r['Segmento']),clean(r['Segmento']))
    if rb and rb not in c['rubros']: c['rubros'].append(rb)
    w=clean(r['Web']).lower()
    if w and w not in c['webs']: c['webs'].append(w)
    if unit and unit not in c['units']: c['units'].append(unit)

S_CLI='Cliente final – operador aéreo / MRO / industria'
S_TECH='Partner tecnológico / competidor'
S_CONS='Canal / aliado – capacitación y consultoría'
S_OEM='Partner OEM – aeronaves, motores y equipos'
S_EPP='Canal / distribuidor EPP y equipamiento táctico'
S_SEG='Aliado complementario – seguridad'
S_SAL='Aliado complementario – salud'
S_GOV='Institución pública – FF.AA., policías y autoridad'
S_INST='Institución / gremio / pabellón país'
S_PROV='Proveedor defensa / industrial'
CUR={
 S_CLI:['LATAM AILINES GROUP','AEROCARDAL LIMITADA','AEROLIDER','Aeroservicios Toqui Ltda','EAGLE COPTERS SOUTH AMERICA','A PROFESIONAL AVIATION SERVICES CHILE ( APAS CHILE)','AEROMAC (AERONAUTICAL MAINTENANCE CENTER SPA.)','NACIONAL DE AERONÁUTICA DE CHILE - ENAER','ASMAR - ASTILLEROS Y MAESTRANZAS DE LA ARMADA','Fabricas Y Maestranzas Del Ejercito','WORLD FUEL SERVICES','FORESTAL ARAUCO S.A','EXPLORAIR','BODE AVIATION','HAIN SPA','ABM AERO','AEREOS DEFENSE','FIELD AEROSPACE','ESMAX','HELISUL TAXI AEREO LTDA','CLUB AÉREO DE SANTIAGO','AERONAUTICAL INVESTMENTS SPA','LAMAIGNERE CARGO S.L.','TKOF SPA','SISDEF LTDA','DISTRIBUIDORA DE VEHICULOS ESPECIALES SPA.'],
 S_TECH:['TELEDYNE FLIR LLC.','THEON SENSORS','TONBO IMAGING INDIA LIMITED','DRONESHIELD','SHIELD AI','AUTEL ROBOTICS','GEOCOM','SOCIEDAD CREA3D INGENIERÍA SPA','ACELERADORA ANIPRIME LIMITADA','ALACRON.AI','ROBOMOTIC SPA','RASTER 4 SPA','ESRI CHILE','PLANET LABS PBC','ICEYE','CAPELLA, AN IONQ COMPANY','SATLANTIS MICROSATS S.A.','ORORATECH','HAWKEYE 360','VOLATUS AEROSPACE CORP','PERSISTENT SYSTEMS','DOODLE LABS. LLC','MOTOROLA SOLUTIONS CHILE','SKYTEL TELECOMUNICACIONES LIMITADA','IPS TELECOM SPA','Codan Communications','NVIS Communications','OPEN CHANNEL S.A.','SETRONIX CHILE S.A','SERVICIOS Y SOLUCIONES TECNOLOGICAS SA','COMERCIALIZADORA RUGGEDTECH LIMITADA','BLIGHTER SURVEILLANCE SYSTEMS LTD.','SI IMAGING SERVICES (SIIS)','DTS (DESARROLLO DE TECNOLOGIA Y SISTEMAS SPA)','TEK CHILE','SMART PARTNERS','RED CAT HOLDING (Teal Drones)','IDEAFORGE TECHNOLOGY LTD.','XMOBOTS AEROESPACIAL E DEFESA','AEVEX','SECUNET INTERNATIONAL GMBH','EDNEO AG','TRAKKA SYSTEMS','JEPPESEN FOREFLIGHT','NEXATLAS TECNOLOGIAS DIGITAIS LTDA.','HUGHES NETWORK SYSTEMS LLC','IBERNEX'],
 S_CONS:['9G SIMULATORS SPA','CORPORACION INSTITUTO PROFESIONAL INACAP','ESCUELA TECNICA AERONAUTICA - DIRECCION GENERAL DE AERONAUTICA CIVIL (ETA)','ADF FLIGHT SCHOOL','EOLO CONSULTORES','REISER SIMULATION & TRAINING GMBH','COMPUTATIONAL ENTERPRISE SIMULATIONS CHILE SPA','EQUIPOS Y SERVICIOS TENSOR LTDA','SERVICIOS PROFESIONALES UP SPA (BLACK UP)','CT INGENIEROS AAI, S.L.'],
 S_OEM:['EMBRAER','AIRBUS SAS','THE BOEING COMPANY','BELL TEXTRON INC.','TEXTRON AVIATION','LEONARDO SPA','CIRRUS AIRCRAFT','DIAMOND AIRCRAFT INDUSTRIES GMBH','DE HAVILLAND AIRCRAFT OF CANADA LIMITED','MD HELICOPTERS','PRATT & WHITNEY','GE AEROSPACE','ROLLS ROYCE - POWER SYSTEMS AG','SAFRAN HELICOPTERS ENGINES BRASIL','THALES','COLLINS AEROSPACE','L3HARRIS TECHNOLOGIES','STANDARDAERO','AIRCRAFT INDUSTRIES','SCHIEBEL ELEKTRONISCHE GERATE GMBH','AIR AMBULANCE TECHNOLOGY GMBH','Comercial Kaufmann S.A','KIA','HYUNDAI','FAUN TRACKWAY LIMITED','ELECTROAIR OU','TECHMAN-HEAD'],
 S_EPP:['COMERCIALIZADORA RMH SPA (HAIX)','BELLEVILLE BOOT COMPANY','INVERSIONES RETA SPA','BTG GROUP SPA','SOCIEDAD COMERCIAL E INDUSTRIAL AUSTRAL TACTICAL SPA','MGCG SPA','PROTACTICAL SPA','CONDOR OUTDOOR PRODUCTS, INC','POINT BLANK ENTERPRISES','MKU LIMITED','MILFORT S.A.','MIGUEL CABALLEROS CHILE SPA','FENIX PROTECTOR S.R.O.','KONG ITALY SPA','SWITLIK PARACHUTE CO., INC.','IMPORTACIONES Y REPRESENTACIONES TRUST LINE LTDA','DON KYATT SOUTHAMERICA SPA','GUMARNY ZUBRI A.S.','ARMOR INTERNATIONAL SAS','IMPORTADORA LILLO SPA'],
 S_SEG:['BLINDATEK SPA','RED FLAME','CHANG SUNG ACE','KFPT. LTD','INER CHILE SPA','HAZTEC INTYERNATIONAL LTD','CAYLYM TECHNOLOGIES INTERNATIONAL','HESCO REP. BY BAI CONSULTING','Tarpulin','S.S TARGET','INSTITUTO DE INVESTIGACIONES Y CONTROL'],
 S_SAL:['WORKMED','CCFA LA ARAUCANA','MUTUALIDAD DEL EJÉRCITO Y AVIACIÓN'],
 S_GOV:['CARABINEROS DE CHILE','FUERZA AÉREA DE CHILE','POLICÍA DE INVESTIGACIONES DE CHILE (PDI)','EJERCITO DE CHILE','Brigada de Aviación del Ejército de Chile','Escuela de Suboficiales del Ejército de Chile','Escuela Militar del Ejército de Chile','DIRECCIÓN GENERAL DE LOS SERVICIOS DE LA ARMADA','DIRECCIÓN GENERAL DE AERONÁUTICA CIVIL (DGAC)','Gendarmeria De Chile','INSTITUTO GEOGRAFICO MILITAR','AGENCIA CHILENA DE EFICIENCIA ENERGETICA (VUELO LIMPIO)','SERVICIO NACIONAL DE TURISMO - SERNATUR','CORPORACION DE LA INDUSTRIA AERONAUTICA COLOMBIANA- CIAC S.A.','DEFENSE SECURITY COOPERATION AGENCY','SSB','MAKINA VE KIMYA ENDUSTRISI (MKE)'],
}
CURK={key(n):s for s,ns in CUR.items() for n in ns}
INSTRX=r'EMBASSY|EMBAJADA|CHAMBER|CAMARA|ASSOCIA|ASOCIACI|AGENCY|AGENCIA|PROMARGENTINA|KOTRA|FERIA|FEINDEF|FAMEX|IDEF|GIFAS|TEDAE|CLUSTER|ADVANTAGE AUSTRIA|BUSINESS FINLAND|MUSEO|FUNDACION|FEDERACION|CIRCULO DE PILOTOS|LIVING HISTORY|PARQUE DE INOVACAO|ECM EXPO|EDEFA|INFORMATION AND DESIGN|ABIMDE|GOBIERNO|COMMERCIAL CORPORATION|ADMINISTRACION Y GESTION DE INSTITUCIONES|SSI\)|AUSTRIAN DEFENCE|ICE AGENCIA|KALLMAN|ACHIDE|POLISH INVESTMENT|CZECH AEROSPACE INDUSTRY'
def seg_of(c):
    k=key(c['Empresa'])
    if k in CURK: return CURK[k]
    n=norm(c['Empresa'])
    if re.search(INSTRX,n): return S_INST
    return S_PROV
ORDER=[S_CLI,S_TECH,S_CONS,S_OEM,S_EPP,S_SEG,S_SAL,S_GOV,S_INST,S_PROV]
PRI={S_CLI:'Alta',S_TECH:'Alta',S_CONS:'Alta',S_OEM:'Media',S_EPP:'Media',S_SEG:'Media',S_SAL:'Media',S_GOV:'Media',S_INST:'Media',S_PROV:'Baja'}
ACC={S_CLI:'Identificar Gerente de Seguridad Operacional/HSE; proponer piloto (hangar, rampa, mantenimiento, combustible)',
 S_TECH:'Mapear oferta: evaluar integración (cámaras térmicas, drones, satélite, comunicaciones) o competencia',
 S_CONS:'Explorar alianza comercial / formación conjunta en seguridad',
 S_OEM:'Explorar alianza con representante local / centros de servicio en Chile',
 S_EPP:'Evaluar como canal de distribución hacia industria/FF.AA.',
 S_SEG:'Evaluar co-venta o integración de soluciones',
 S_SAL:'Evaluar alianza en salud ocupacional / fatiga',
 S_GOV:'Mapear unidad de prevención/seguridad operacional; vía licitación o convenio marco',
 S_INST:'Networking institucional; identificar contacto',
 S_PROV:'Nutrir con contenido; sin acción comercial inmediata'}
out_rows=[]
for k,c in comp.items():
    seg=seg_of(c); pri=PRI[seg]
    if seg==S_PROV and c['pais']=='Chile': pass
    notas="FIDAE 2026."
    if c['units']: notas+=f" Unidades presentes: {len(c['units'])+1}."
    if len(c['stands'])>1: notas+=" Varios espacios (mayor inversión)."
    if not c['stands']: notas+=" Sin ubicación asignada."
    cruce=[]; d=None
    for nm,m in CRMS:
        x=m.get(k)
        if x:
            cruce.append(f"{nm} (ID {x['ID']})")
            if nm=='Directorio SSO': d=x
    if cruce: notas+=" También en: "+'; '.join(cruce)+"."
    if d and c['pais']!='Perú': notas+=" Contacto heredado = filial Perú."
    if not d: notas+=" Sin contacto publicado."
    src_=d or {}
    out_rows.append({'Nombre':src_.get('Nombre') or '','Cargo':src_.get('Cargo') or '','Empresa':c['Empresa'],'País':c['pais'],
      'Email':src_.get('Email') or '','LinkedIn':'','Origen':'FIDAE 2026','Estado':'Prospecto','Prioridad':pri,
      'Próxima acción':ACC[seg],'Notas':notas,'Segmento SafetyMind':seg,'Stand(s)':' | '.join(c['stands']),
      'N° stands':len(c['stands']),'Rubro FIDAE':' | '.join(c['rubros']),'Web':'; '.join(c['webs']),
      'Teléfono(s)':src_.get('Teléfono(s)') or '','Cruce otros CRM':'; '.join(cruce) or 'No'})
porder={'Alta':0,'Media':1,'Baja':2}
out_rows.sort(key=lambda r:(porder[r['Prioridad']],ORDER.index(r['Segmento SafetyMind']),r['País']!='Chile',norm(r['Empresa'])))
COLS=['ID','Nombre','Cargo','Empresa','País','Email','LinkedIn','Origen','Estado','Prioridad','Próxima acción','Notas','Segmento SafetyMind','Stand(s)','N° stands','Rubro FIDAE','Web','Teléfono(s)','Cruce otros CRM']
wb=Workbook(); HF=PatternFill('solid',fgColor='1F3A5F'); HFont=Font(bold=True,color='FFFFFF')
def header(ws,cols,widths):
    ws.append(cols)
    for i,cn in enumerate(cols,1):
        cell=ws.cell(1,i); cell.fill=HF; cell.font=HFont; cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1] if i-1<len(widths) else 16
    ws.freeze_panes='A2'
ws=wb.active; ws.title='Contactos'
header(ws,COLS,[6,20,20,45,16,30,14,14,12,10,50,50,45,22,9,30,32,24,30])
for i,r in enumerate(out_rows,1):
    ws.append([i]+[r[c] for c in COLS[1:]]); ws.cell(ws.max_row,18).number_format='@'
n=len(out_rows)+1
t=Table(displayName='Contactos',ref=f"A1:{get_column_letter(len(COLS))}{n}"); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
dv1=DataValidation(type='list',formula1='"Prospecto,Contactado,Reunión,Propuesta,Cliente,No interesado"',allow_blank=True); dv1.add(f'I2:I{n+500}')
dv2=DataValidation(type='list',formula1='"Alta,Media,Baja,Por calificar"',allow_blank=True); dv2.add(f'J2:J{n+500}')
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
ws2=wb.create_sheet('Oportunidades'); header(ws2,['ID oportunidad','ID contacto','Empresa','Necesidad','Solución propuesta','Etapa','Valor estimado','Probabilidad','Próxima acción','Fecha'],[14,12,40,40,40,14,14,12,40,12])
ws3=wb.create_sheet('Resumen'); header(ws3,['Evento/Fuente','País','Cantidad','',''],[50,26,12,12,12])
pc=collections.Counter(r['País'] for r in out_rows)
for i,(p,_) in enumerate(pc.most_common(),2): ws3.append(['FIDAE 2026',p,f'=COUNTIFS(Contactos!$H:$H,$A{i},Contactos!$E:$E,$B{i})'])
last=ws3.max_row; ws3.append(['Total','',f'=SUM(C2:C{last})']); ws3.append([''])
def hdr(row,k):
    for cc in range(1,k+1): ws3.cell(row,cc).fill=HF; ws3.cell(row,cc).font=HFont
r0=ws3.max_row+1; ws3.append(['Segmento SafetyMind','Alta','Media','Baja','Total']); hdr(r0,5)
for s in ORDER:
    i=ws3.max_row+1; ws3.append([s]+[f'=COUNTIFS(Contactos!$M:$M,$A{i},Contactos!$J:$J,{col}${r0})' for col in 'BCD']+[f'=SUM(B{i}:D{i})'])
i=ws3.max_row+1; ws3.append(['Total']+[f'=SUM({cl}{r0+1}:{cl}{i-1})' for cl in 'BCDE']); ws3.append([''])
r1=ws3.max_row+1; ws3.append(['Estado (pipeline)','Cantidad']); hdr(r1,2)
for e_ in ['Prospecto','Contactado','Reunión','Propuesta','Cliente','No interesado']:
    i=ws3.max_row+1; ws3.append([e_,f'=COUNTIF(Contactos!$I:$I,A{i})'])
ws3.append(['']); r2=ws3.max_row+1; ws3.append(['Indicadores','Cantidad']); hdr(r2,2)
for lab,f in [('Registros originales',len(rows)),('Empresas / instituciones únicas','=COUNTA(Contactos!$D$2:$D$3000)'),('Con sede en Chile','=COUNTIF(Contactos!$E$2:$E$3000,"Chile")'),('Con web','=COUNTIF(Contactos!$Q$2:$Q$3000,"?*")'),('En otros CRM','=COUNTIF(Contactos!$S$2:$S$3000,"*(ID*")')]:
    ws3.append([lab,f])
ws4=wb.create_sheet('Eventos'); header(ws4,['Evento','País','Valor para SafetyMind'],[45,12,100])
for e in [['Directorio SSO 2026 – Cero Accidentes (22.ª ed.)','Perú','Mapa de ~670 proveedores SST en Perú (ver CRM_Directorio_SSO_2026).'],
 ['ExpoMina Perú 2026','Perú','Feria minera (ver CRM_ExpoMina_2026).'],['Exponor 2026 (Antofagasta)','Chile','Feria minera del norte de Chile (ver CRM_Exponor_2026).'],
 ['AquaSur 2026 (Puerto Montt)','Chile','Feria acuícola (ver CRM_AquaSur_2026).'],
 ['FIDAE 2026 (Santiago)','Chile','Feria aeroespacial y de defensa: operadores aéreos y MRO (hangar, rampa, combustible), FF.AA. y policías, tecnología de sensores/drones/satélites y fabricantes de aeronaves.']]: ws4.append(e)
wb.save(out)
print(len(rows),len(out_rows)); print(collections.Counter(r['Segmento SafetyMind'] for r in out_rows)); print(collections.Counter(r['Prioridad'] for r in out_rows))
print('cruce',sum(r['Cruce otros CRM']!='No' for r in out_rows))
print('not found',[x for ns in CUR.values() for x in ns if key(x) not in comp])
json.dump(out_rows,open(out+'.json','w'),ensure_ascii=False)
