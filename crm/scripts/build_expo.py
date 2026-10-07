import csv, re, sys, json, collections, unicodedata
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
src, dirxlsx, out = sys.argv[1:4]
def fixmoji(s):
    for a,b in [('ÃƑÂ€˜','Ñ'),('ÃƑÂ\x8d','Í'),('ÃƑÂ','Í'),('Ã\x8d','Í'),('Ã“','Ó'),('Ã‰','É'),('ï¼Œ',','),('Â',' '),('\xa0',' '),('PERÊ','PERÚ'),(' €Œ',''),('Ï¼Œ',','),(' €™',"'")]: s=s.replace(a,b)
    for _ in range(3):
        if 'Ã' not in s and 'Â' not in s: break
        try: s=s.encode('cp1252').decode('utf8')
        except Exception:
            try: s=s.encode('latin1').decode('utf8')
            except Exception: break
    return s.replace('PERÊ','PERÚ')
def clean(s): return re.sub(r'\s+',' ',fixmoji(s)).strip()
def norm(s): return unicodedata.normalize('NFKD',s.upper()).encode('ascii','ignore').decode()
def key(n):
    n=norm(n).lower()
    n=re.split(r'\s+-\s+',n)[0]
    n=re.sub(r'\b(s\.?\s?a\.?\s?c\.?|s\.?\s?a\.?\s?a\.?|s\.?\s?a\.?|e\.?i\.?r\.?l\.?|s\.?r\.?l\.?|s\.?c\.?r\.?l\.?|ltda\.?|sac|eirl|srl|peru|del)\b','',n)
    return re.sub(r'[^a-z0-9]','',n)
# Directorio CRM
dwb=openpyxl.load_workbook(dirxlsx); dws=dwb['Contactos']
dh=[c.value for c in dws[1]]
dirmap={}
for r in dws.iter_rows(min_row=2,values_only=True):
    d=dict(zip(dh,r))
    if not d.get('Empresa'): continue
    dirmap.setdefault(key(d['Empresa']),d)
rows=list(csv.DictReader(open(src,encoding='utf-8-sig')))
comp=collections.OrderedDict()
for r in rows:
    name=clean(r['Empresa']); k=key(name)
    c=comp.setdefault(k,{'Empresa':name,'stands':[],'webs':[]})
    st=clean(r['Stand']); w=clean(r['Web']).lower()
    if st and st not in c['stands']: c['stands'].append(st)
    if w and w!='-' and w not in c['webs']: c['webs'].append(w)
CHN=r'SHANDONG|ZHEJIANG|JIANGSU|HEBEI|HENAN|ZHENGZHOU|SHANGHAI|BEIJING|GUANGDONG|SHENYANG|TIANJIN|QINGDAO|WUXI|NANTONG|CHANGZHOU|JIAOZUO|BAOJI|TAIAN|LUOYANG|HUNAN|HUBEI|ANHUI|FUJIAN|SHAANXI|SHANXI|XIAMEN|NINGBO|SUZHOU|HANGZHOU|YANTAI|WEIFANG|JINAN|CHANGSHA|XUZHOU|SICHUAN|CHENGDU|LIAONING|DALIAN|YUNNAN|JIANGXI|GANSU|INNER MONGOLIA|SHENZHEN|GUANGZHOU|CHINA|SINO'
TLD=[(r'\.com\.pe|\.pe(/|$)','Perú'),(r'\.cl(/|$)','Chile'),(r'\.cn(/|$)','China'),(r'\.com\.br|\.br(/|$)','Brasil'),(r'\.de(/|$)','Alemania'),(r'\.ca(/|$)','Canadá'),(r'\.com\.au|\.au(/|$)','Australia'),(r'\.it(/|$)','Italia'),(r'\.es(/|$)','España'),(r'\.mx(/|$)','México'),(r'\.co(/|$)','Colombia'),(r'\.com\.ar|\.ar(/|$)','Argentina'),(r'\.fi(/|$)','Finlandia'),(r'\.se(/|$)','Suecia'),(r'\.in(/|$)','India'),(r'\.tr(/|$)','Turquía'),(r'\.za(/|$)','Sudáfrica'),(r'\.uk(/|$)','Reino Unido'),(r'\.at(/|$)','Austria'),(r'\.ch(/|$)','Suiza'),(r'\.kr(/|$)','Corea del Sur'),(r'\.jp(/|$)','Japón'),(r'\.ru(/|$)','Rusia'),(r'\.pl(/|$)','Polonia'),(r'\.nl(/|$)','Países Bajos'),(r'\.be(/|$)','Bélgica'),(r'\.fr(/|$)','Francia')]
def pais(name,webs,d):
    n=norm(name); w=' '.join(webs)
    m=re.search(r'EMBAJADA DE ([A-Z ]+)|PRO ?CHILE',n)
    if 'PROCHILE' in n.replace(' ',''): return 'Chile'
    if re.search(r'PERU|S\.?\s?A\.?\s?C\.?|E\.?\s?I\.?\s?R\.?\s?L\.?|\bSAC\b|\bEIRL\b|S\.?\s?R\.?\s?L\.?|S\.?\s?A\.?\s?A\.?',n): return 'Perú'
    if re.search(CHN,n) or re.search(r'CO\.?,? ?LTD',n): return 'China'
    for kw,pp in [('MEXICO','México'),('CHILE','Chile'),('INDIA','India'),('SOUTH AFRICA','Sudáfrica'),('BRASIL','Brasil'),('COLOMBIA','Colombia'),('ARGENTINA','Argentina'),('ESTADOS UNIDOS','Estados Unidos'),('CANADA','Canadá'),('ALEMANA','Alemania')]:
        if kw in n: return pp
    if 'PVT' in n: return 'India'
    if 'GMBH' in n: return 'Alemania'
    if 'BRASIL' in n or 'LTDA' in n and '.br' in w: return 'Brasil'
    if re.search(r'\bOY\b',n): return 'Finlandia'
    if re.search(r'\bAB\b',n): return 'Suecia'
    for pat,p in TLD:
        if re.search(pat,w): return p
    if re.search(r'\bS\.?P\.?A\.?\b',n): return 'Chile (por validar)'
    if re.search(r'QUELLAVECO|CERRO VERDE',n): return 'Perú'
    if re.search(r'SOCIEDAD ANONIMA|\bS\.? ?A\.?$|\bS A\b',n): return 'Perú (por validar)'
    if re.search(r'\bINC\b|\bLLC\b|\bCORP\b|\bLP\b',n): return 'Estados Unidos (por validar)'
    return 'Por validar'
OEM=r'FERREYROS|KOMATSU|EPIROC|SANDVIK|CATERPILLAR|LIEBHERR|NORMET|RESEMIN|HITACHI|XCMG|SANY|ZOOMLION|VOLVO|SCANIA|DIVEMOTOR|BELAZ|METSO|FLSMIDTH|WEIR|ATLAS COPCO|BOART|DERCO|UNIMAQ|MOTORED|MARCO PERUANA|SIEMAG|TERNIUM|GOODYEAR|BRIDGESTONE|MICHELIN|TITAN'
TECH=r'HEXAGON|MAPTEK|CAD SOLUTION|SOUTHERN TECHNOLOGY|YOKOGAWA|DATAMINE|DEUTSCHE TELEKOM|DIGITAL|SOFTWARE|\bDATA\b|ANALYTIC|AUTOMATI|TELECOM|\bENTEL\b|INTEGRATEL|CLARO|ROBOT|DRON|\bIA\b|\bAI\b|VISION|CAMERA|\bIOT\b|CLOUD|SATELIT|SATELLIT|WENCO|MODULAR MINING|TELEMETR|MONITOREO|SMART|INTELIGEN|RADIO\b|COMUNICACION|COMMUNICATION|CONECTIV|TESACOM|SIEMENS|ABB\b|ROCKWELL|SCHNEIDER|HONEYWELL|GEOSYSTEM|TRIMBLE|SENSOR'
MINE=r'COMPA[NÑ]IA MINERA|CIA\.? MINERA|SOCIEDAD MINERA|MINERA [A-Z]+ S\.?A|MINING (COMPANY|CORP)|SOUTHERN PERU|ANTAMINA|ANGLO AMERICAN|FREEPORT|CERRO VERDE|BUENAVENTURA|VOLCAN|HOCHSCHILD|NEXA|MMG|LAS BAMBAS|CHINALCO|SHOUGANG|MARSA|PODEROSA|GOLD FIELDS|NEWMONT|YANACOCHA|BARRICK|GLENCORE|ANTAPACCAY|\bTECK\b|ACTIVOS MINEROS|MINSUR|RAURA|\bARES\b|BOROO|PAN AMERICAN|SIERRA METALS|BEAR CREEK|ZIJIN|CODELCO'
CONTR=r'CONTRATIST|MINERIA SUBTERRANEA|COMISUB|DRILLING|PERFORA|EXPLOSIV|VOLADUR|SERVICIOS MINEROS|MINING SERVICES|CONSTRUC|INGENIERIA Y CONSTRUC|TRANSPORTES (LINEA|MERIDIAN|POLUX)|MOVIMIENTO DE TIERRA|MANTENIMIENTO|MONTAJE|STRACON|COSAPI|SAN MARTIN|MOTA-ENGIL|JJC|GYM\b|CONFIPETROL|EXPLOMIN|ROBOCON|SEMPER'
CONS=r'CONSULT|CAPACITA|ASESOR|CERTIFIC|INSPECC|\bSGS\b|BUREAU VERITAS|TUV|INTERTEK|INSTITUTO|UNIVERSIDAD|TECSUP|CETEMIN|ISEM|AUDIT'
EPP=r'SAFETY|SEGURIDAD INDUSTRIAL|\bEPP\b|PROTECCION PERSONAL|GUANTE|CALZADO|BOTAS|\b3M\b|MSA\b|DRAEGER|DRAGER|ANDES SEGURIDAD|SEGURINDUSTRIA|STEELPRO|ROPA|UNIFORM|INDUMENTARIA'
SEG=r'SEGURIDAD|FIRE|INCENDIO|VIGILANCIA|RESCATE|SE[NÑ]ALIZ|DETECCION DE GAS|GAS DETECT|BLINDAD|EXTINT|ANTICOLISION|FATIG'
SAL=r'SALUD|MEDIC|CLINICA|OCUPACIONAL|LABORATORIO CLIN'
INST=r'EMBAJADA|CAMARA|PROCHILE|ASOCIACION|MINISTERIO|SNMPE|GREMIO|CONSEJO|PROINVERSION|INGEMMET|OSINERGMIN|SUNAFIL|GOBIERNO|TRADE COMMISSION|PABELLON|PAVILION|INVEST'
SEGS0=[('Minera (cliente final)',MINE,'Alta','Mapear Gerente SSOMA/HSE y unidades mineras; proponer piloto'),
 ('Cliente potencial – contratista / operación de alto riesgo',CONTR,'Alta','Identificar Gerente SSOMA/HSE en LinkedIn y proponer demo'),
 ('Partner tecnológico / competidor',TECH,'Alta','Mapear oferta: evaluar integración o competencia'),
 ('Canal / aliado – consultoría SST',CONS,'Alta','Explorar alianza comercial / reventa a sus clientes'),
 ('Partner OEM – equipos mineros',OEM,'Media','Explorar integración con flota/equipos (visión + telemetría)'),
 ('Canal / distribuidor EPP',EPP,'Media','Evaluar como canal de distribución hacia minería/industria'),
 ('Aliado complementario – seguridad',SEG,'Media','Evaluar co-venta o integración de soluciones'),
 ('Aliado complementario – salud',SAL,'Media','Evaluar alianza en salud ocupacional / fatiga'),
 ('Institución / gremio / pabellón país',INST,'Media','Networking institucional; identificar contacto de pabellón')]
_o=['Institución / gremio / pabellón país','Minera (cliente final)','Canal / distribuidor EPP','Aliado complementario – seguridad','Cliente potencial – contratista / operación de alto riesgo','Partner OEM – equipos mineros','Partner tecnológico / competidor','Canal / aliado – consultoría SST','Aliado complementario – salud']
SEGS=sorted(SEGS0,key=lambda x:_o.index(x[0]))
DSEG_MAP={'Cliente potencial – operación de alto riesgo':'Cliente potencial – contratista / operación de alto riesgo'}
ORDER=[s[0] for s in SEGS0]+['Proveedor industrial (equipos/insumos)']
ACC={s[0]:s[3] for s in SEGS}; ACC['Proveedor industrial (equipos/insumos)']='Nutrir con contenido; sin acción comercial inmediata'
PRI={s[0]:s[2] for s in SEGS}; PRI['Proveedor industrial (equipos/insumos)']='Baja'
out_rows=[]
for k,c in comp.items():
    n=norm(c['Empresa'])+' '+' '.join(c['webs']).upper()
    seg=None
    d=dirmap.get(k)
    p=pais(c['Empresa'],c['webs'],d)
    for s,pat,_,_ in SEGS:
        if re.search(pat,n):
            if p=='China' and s in ('Cliente potencial – contratista / operación de alto riesgo','Partner OEM – equipos mineros','Minera (cliente final)'): continue
            seg=s; break
    if p=='China' and seg=='Partner tecnológico / competidor' and not re.search(r'DIGITAL|SOFTWARE|DATA|AI\b|VISION|CAMERA|IOT|CLOUD|ROBOT',n): seg=None
    if seg is None and d:
        ds=DSEG_MAP.get(d['Segmento SafetyMind'],d['Segmento SafetyMind'])
        if ds in ORDER: seg=ds
    if seg is None: seg='Proveedor industrial (equipos/insumos)'
    pri=PRI[seg]
    if d and d['Anunciante']=='Sí' and pri!='Alta': pri={'Media':'Alta','Baja':'Media'}[pri]
    if p in ('Por validar',) and d: p=d['País']
    notas="ExpoMina 2026."
    if len(c['stands'])>1: notas+=" Varios stands (mayor inversión)."
    if d: notas+=f" También en CRM Directorio SSO 2026 (ID {d['ID']})."
    if '(por validar)' in p or p=='Por validar': notas+=" País inferido; validar."
    if not d: notas+=" Sin contacto publicado."
    out_rows.append({'Nombre':(d or {}).get('Nombre') or '','Cargo':(d or {}).get('Cargo') or '','Empresa':c['Empresa'],'País':p,
      'Email':(d or {}).get('Email') or '','LinkedIn':'','Origen':'ExpoMina 2026','Estado':'Prospecto','Prioridad':pri,
      'Próxima acción':ACC[seg],'Notas':notas,'Segmento SafetyMind':seg,'Stand(s)':', '.join(c['stands']),'N° stands':len(c['stands']),
      'Web':'; '.join(c['webs']),'Teléfono(s)':(d or {}).get('Teléfono(s)') or '','Dirección':(d or {}).get('Dirección') or '',
      'En Directorio SSO':'Sí' if d else 'No'})
porder={'Alta':0,'Media':1,'Baja':2}
out_rows.sort(key=lambda r:(porder[r['Prioridad']],ORDER.index(r['Segmento SafetyMind']),-r['N° stands'],r['Empresa']))
COLS=['ID','Nombre','Cargo','Empresa','País','Email','LinkedIn','Origen','Estado','Prioridad','Próxima acción','Notas','Segmento SafetyMind','Stand(s)','N° stands','Web','Teléfono(s)','Dirección','En Directorio SSO']
wb=Workbook(); HF=PatternFill('solid',fgColor='1F3A5F'); HFont=Font(bold=True,color='FFFFFF')
def header(ws,cols,widths):
    ws.append(cols)
    for i,cn in enumerate(cols,1):
        cell=ws.cell(1,i); cell.fill=HF; cell.font=HFont; cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1] if i-1<len(widths) else 16
    ws.freeze_panes='A2'
ws=wb.active; ws.title='Contactos'
header(ws,COLS,[6,20,20,45,18,30,14,16,12,10,45,55,40,18,9,32,28,45,12])
for i,r in enumerate(out_rows,1):
    ws.append([i]+[r[c] for c in COLS[1:]]); ws.cell(ws.max_row,17).number_format='@'
n=len(out_rows)+1
t=Table(displayName='Contactos',ref=f"A1:{get_column_letter(len(COLS))}{n}"); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
dv1=DataValidation(type='list',formula1='"Prospecto,Contactado,Reunión,Propuesta,Cliente,No interesado"',allow_blank=True); dv1.add(f'I2:I{n+500}')
dv2=DataValidation(type='list',formula1='"Alta,Media,Baja,Por calificar"',allow_blank=True); dv2.add(f'J2:J{n+500}')
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
ws2=wb.create_sheet('Oportunidades'); header(ws2,['ID oportunidad','ID contacto','Empresa','Necesidad','Solución propuesta','Etapa','Valor estimado','Probabilidad','Próxima acción','Fecha'],[14,12,40,40,40,14,14,12,40,12])
ws3=wb.create_sheet('Resumen'); header(ws3,['Evento/Fuente','País','Cantidad','',''],[50,26,12,12,12])
pc=collections.Counter(r['País'] for r in out_rows)
for i,(p,_) in enumerate(pc.most_common(),2): ws3.append(['ExpoMina 2026',p,f'=COUNTIFS(Contactos!$H:$H,$A{i},Contactos!$E:$E,$B{i})'])
last=ws3.max_row; ws3.append(['Total','',f'=SUM(C2:C{last})']); ws3.append([''])
def hdr(row,k):
    for cc in range(1,k+1): ws3.cell(row,cc).fill=HF; ws3.cell(row,cc).font=HFont
r0=ws3.max_row+1; ws3.append(['Segmento SafetyMind','Alta','Media','Baja','Total']); hdr(r0,5)
for s in ORDER:
    i=ws3.max_row+1; ws3.append([s]+[f'=COUNTIFS(Contactos!$M:$M,$A{i},Contactos!$J:$J,{col}${r0})' for col in 'BCD']+[f'=SUM(B{i}:D{i})'])
i=ws3.max_row+1; ws3.append(['Total']+[f'=SUM({cl}{r0+1}:{cl}{i-1})' for cl in 'BCDE']); ws3.append([''])
r1=ws3.max_row+1; ws3.append(['Estado (pipeline)','Cantidad']); hdr(r1,2)
for e in ['Prospecto','Contactado','Reunión','Propuesta','Cliente','No interesado']:
    i=ws3.max_row+1; ws3.append([e,f'=COUNTIF(Contactos!$I:$I,A{i})'])
ws3.append(['']); r2=ws3.max_row+1; ws3.append(['Indicadores','Cantidad']); hdr(r2,2)
for lab,f in [('Registros originales (stands)',len(rows)),('Empresas únicas','=COUNTA(Contactos!$D$2:$D$3000)'),('Con web','=COUNTIF(Contactos!$P$2:$P$3000,"?*")'),('También en Directorio SSO','=COUNTIF(Contactos!$S:$S,"Sí")'),('Con email (vía Directorio)','=COUNTIF(Contactos!$F$2:$F$3000,"?*")'),('Con 2+ stands','=COUNTIF(Contactos!$O$2:$O$3000,">1")')]:
    ws3.append([lab,f])
ws4=wb.create_sheet('Eventos'); header(ws4,['Evento','País','Valor para SafetyMind'],[45,12,100])
ws4.append(['Directorio SSO 2026 – Cero Accidentes (22.ª ed.)','Perú','Mapa de ~670 proveedores SST en Perú (ver CRM_Directorio_SSO_2026).'])
ws4.append(['ExpoMina Perú 2026','Perú','Feria minera: mineras, contratistas, OEM de equipos, tecnología y proveedores; base para agendar reuniones en stand y prospección post-evento.'])
wb.save(out)
print(len(rows),len(out_rows)); print(collections.Counter(r['Segmento SafetyMind'] for r in out_rows)); print(collections.Counter(r['Prioridad'] for r in out_rows)); print(pc.most_common(15)); print(sum(r['En Directorio SSO']=='Sí' for r in out_rows))
json.dump(out_rows,open(out+'.json','w'),ensure_ascii=False)
