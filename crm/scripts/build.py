import json, re, sys, collections, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
d=json.load(open(sys.argv[1]))
MAP={'Ceneris':'Ceneris E.I.R.L.','Gisber':'Gisbert Representantes E.I.R.L.','Segurindustria':'Segurindustria S.A.','Andes-safety':'Andes Seguridad S.A.C.','Ifsec-la':'IFSEC Perú S.A.C.','Seprocal':'Operaciones Seprocal S.A.C.','Safetymanagment':'Safety Management Resources S.R.L.','Safetymanagement':'Safety Management Resources S.R.L.','Www':'Safety Management Resources S.R.L.','Management':'CleanSpace Respirators (Safety Management Resources)','Safestartlatam':'SafeStart','Coelpra':'Coelpra S.A.C.','Fedseg':'FEDSEG S.A.C.','Csgperu':'CSG Perú S.A.C.','Moldex':'Moldex-Metric Inc.','Ophtalmiccenter':'Ophtalmic Center E.I.R.L.','Foxperu':'Rafael Dávila / Confecciones Fox','Ycare':'Safe & Care S.A.C.','Visionone':'Vision One International S.R.L.','Carp':'Carp & Asociados S.C.R.L.','Pulsosalud':'Pulso Salud S.A.C.','Vixora':'Vixora S.A.','Castem':'Castem E.I.R.L.'}
PAIS={'Moldex-Metric Inc.':'Estados Unidos','SafeStart':'Brasil','CleanSpace Respirators (Safety Management Resources)':'Australia'}
SEG={ # rubro -> (segmento, prioridad base)
 'Empresas contratistas y conexas':('Cliente potencial – operación de alto riesgo',3),
 'Perforación geológica':('Cliente potencial – operación de alto riesgo',3),
 'Voladura, explosivos y polvorines':('Cliente potencial – operación de alto riesgo',3),
 'Sostenimiento de rocas y suelos':('Cliente potencial – operación de alto riesgo',3),
 'Transporte de combustibles y materiales peligrosos':('Cliente potencial – operación de alto riesgo',3),
 'Mantenimiento industrial':('Cliente potencial – operación de alto riesgo',3),
 'Capacitación, consultoría y auditoría':('Canal / aliado – consultoría SST',3),
 'Certificación, inspección y peritaje':('Canal / aliado – consultoría SST',3),
 'Tecnología de la información, comunicación y telecomunicaciones':('Partner tecnológico / competidor',3),
 'Seguridad vial y prevención de accidentes vehiculares':('Aliado complementario – seguridad',2),
 'Detección de gases: equipos y servicios':('Aliado complementario – seguridad',2),
 'Seguridad y vigilancia interna':('Aliado complementario – seguridad',2),
 'Protección contra incendios':('Aliado complementario – seguridad',2),
 'Rescate y salvamento':('Aliado complementario – seguridad',2),
 'Andamios y trabajos en altura':('Aliado complementario – seguridad',2),
 'Equipos de protección personal':('Canal / distribuidor EPP',2),
 'Ropa industrial':('Canal / distribuidor EPP',2),
 'Reflectivos y señalización':('Canal / distribuidor EPP',2),
 'Salud ocupacional':('Aliado complementario – salud',2),
}
ORDER=['Cliente potencial – operación de alto riesgo','Partner tecnológico / competidor','Canal / aliado – consultoría SST','Aliado complementario – seguridad','Canal / distribuidor EPP','Aliado complementario – salud','Proveedor industrial (baja afinidad)']
ACC={'Cliente potencial – operación de alto riesgo':'Identificar Gerente SSOMA/HSE en LinkedIn y proponer demo',
 'Partner tecnológico / competidor':'Mapear oferta: evaluar integración o competencia',
 'Canal / aliado – consultoría SST':'Explorar alianza comercial / reventa a sus clientes',
 'Aliado complementario – seguridad':'Evaluar co-venta o integración de soluciones',
 'Canal / distribuidor EPP':'Evaluar como canal de distribución hacia minería/industria',
 'Aliado complementario – salud':'Evaluar alianza en salud ocupacional / fatiga',
 'Proveedor industrial (baja afinidad)':'Nutrir con contenido; sin acción comercial inmediata'}
def key(n):
    n=unicodedata.normalize('NFKD',n.lower()).encode('ascii','ignore').decode()
    n=re.sub(r'\b(s\.?a\.?c\.?|s\.?a\.?|e\.?i\.?r\.?l\.?|s\.?r\.?l\.?|s\.?c\.?r\.?l\.?|ltda\.?|sac|eirl|srl)\b','',n)
    return re.sub(r'[^a-z0-9]','',n)
def split(s,sep): return [x.strip() for x in s.split(sep) if x.strip()] if s else []
for e in d:
    if e['pagina']==7 and e['rubro'].startswith('Tecnolog'): e['rubro']='Andamios y trabajos en altura'
comp=collections.OrderedDict()
for e in d:
    name=e['empresa']
    anun=False
    if e['orphan']:
        base=name.replace(' (nombre por validar)','')
        name=MAP.get(base,name); anun=base in MAP
    if not e['orphan'] and not (e['email'] or e['telefonos'] or e['web']) and not e['direccion']: continue
    k=key(name)
    c=comp.setdefault(k,{'empresa':name,'rubros':[],'emails':[],'tels':[],'webs':[],'dir':'','ciudad':'','contactos':[],'paginas':[],'anun':False})
    if e['rubro'] not in c['rubros']: c['rubros'].append(e['rubro'])
    for x in split(e['email'].lower(),';'):
        if x.lower() not in [y.lower() for y in c['emails']]: c['emails'].append(x)
    for x in split(e['telefonos'],' / '):
        xn=re.sub(r'\D','',x)[-9:]
        if xn not in [re.sub(r'\D','',y)[-9:] for y in c['tels']]: c['tels'].append(x)
    for x in split(e['web'],';'):
        if x.lower() not in [y.lower() for y in c['webs']]: c['webs'].append(x)
    e['direccion']=re.sub(r'\S*@\S*(\s\.?\w+\b)?','',e['direccion']).strip()
    e['email']=e['email'].lower()
    if len(e['direccion'])>len(c['dir']): c['dir']=e['direccion']; c['ciudad']=e['ciudad']
    if e['contacto'] and e['contacto'] not in c['contactos']: c['contactos'].append(e['contacto'])
    if e['pagina'] not in c['paginas']: c['paginas'].append(e['pagina'])
    c['anun']=c['anun'] or anun
# names of advertisers that appear also with proper names
ANUN_KEYS={key(v) for v in MAP.values()}
FIX={'Ophtalmic Center E.I.R.L.':'ventas@ophtalmiccenter.com','Pulso Salud S.A.C.':'informes@pulsosalud.com'}
for c in comp.values():
    if c['empresa'] in FIX and not c['emails']: c['emails']=[FIX[c['empresa']]]
    c['dir']=re.sub(r'\s{2,}',' ',re.sub(r'\s+fi\b','',c['dir'])).strip(' ,')
    c['tels']=[t for t in c['tels'] if len(re.sub(r'\D','',t))>=6 and not re.fullmatch(r'\d{5}-\d{3}',t)]
rows=[]
for c in comp.values():
    segs=[SEG.get(r,('Proveedor industrial (baja afinidad)',1)) for r in c['rubros']]
    CORE={'Empresas contratistas y conexas','Perforación geológica','Voladura, explosivos y polvorines','Sostenimiento de rocas y suelos'}
    if len(c['rubros'])>=4 and not (CORE & set(c['rubros'])):
        # proveedor multi-rubro: segmento dominante (excluye cliente potencial)
        cand=[x for x in segs if not x[0].startswith('Cliente') and not x[0].startswith('Proveedor')] or segs
        cnt=collections.Counter(x[0] for x in cand)
        top=max(cnt.values()); best=[x for x in cand if cnt[x[0]]==top]
        seg=min(best,key=lambda s:(-s[1],ORDER.index(s[0])))
    else:
        seg=min(segs,key=lambda s:(-s[1],ORDER.index(s[0])))
    score=seg[1]+(1 if (c['anun'] or key(c['empresa']) in ANUN_KEYS) else 0)
    pri={4:'Alta',3:'Alta',2:'Media',1:'Baja'}[score]
    nombre=cargo=''
    if c['contactos']:
        m=re.match(r'^(.*?)\s*\((.*?)\)?$',c['contactos'][0])
        if m and '(' in c['contactos'][0]: nombre,cargo=m.group(1).strip(),m.group(2).strip()
        else: nombre=c['contactos'][0]
        if nombre.lower().startswith('oficina'): nombre=''
        nombre=re.sub(r'^comercial en Perú:\s*','',nombre)
    pais=PAIS.get(c['empresa'],'Perú')
    if re.search(r'Colombia$',c['dir']): pais='Colombia'
    elif re.search(r'Paraná|Brasil$',c['dir']): pais='Brasil'
    elif 'Argentina' in c['dir'] and 'CABA' in c['dir']: pais='Argentina'
    elif re.search(r'China$',c['dir']): pais='China'
    elif re.search(r'Canadá',c['dir']) and 'Lima' not in c['dir']: pais='Canadá'
    anun=c['anun'] or key(c['empresa']) in ANUN_KEYS
    notas=f"Directorio SSO 2026, pág. PDF {', '.join(map(str,sorted(c['paginas'])))}."
    if anun: notas+=" Anunciante (inversión activa en SST)."
    if nombre: notas+=" Contacto publicado; validar vigencia."
    else: notas+=" Sin contacto nominal; no inventar decisor."
    if len(c['emails'])>1: notas+=" Otros emails: "+'; '.join(c['emails'][1:])+'.'
    rows.append({'Nombre':nombre,'Cargo':cargo,'Empresa':c['empresa'],'País':pais,'Email':c['emails'][0] if c['emails'] else '',
      'LinkedIn':'','Origen':'Directorio SSO 2026','Estado':'Prospecto','Prioridad':pri,'Próxima acción':ACC[seg[0]],'Notas':notas,
      'Rubro(s)':'; '.join(c['rubros']),'Segmento SafetyMind':seg[0],'Teléfono(s)':' / '.join(c['tels']),'Web':'; '.join(c['webs']),
      'Dirección':c['dir'],'Ciudad/Región':c['ciudad'],'Anunciante':'Sí' if anun else 'No'})
FIXES=json.load(open(sys.argv[3])) if len(sys.argv)>3 else {}
for r in rows:
    for k,v in FIXES.get(r['Empresa'],{}).items(): r[k]=v
    r['Teléfono(s)']=' / '.join(t for t in r['Teléfono(s)'].split(' / ') if t and not re.fullmatch(r'\d{2,4}-\d{2,4}',t))
    r['Dirección']=re.sub(r'Mz\.\.','Mz.',r['Dirección'])
porder={'Alta':0,'Media':1,'Baja':2}
rows.sort(key=lambda r:(porder[r['Prioridad']],ORDER.index(r['Segmento SafetyMind']),r['Empresa'].lower()))
wb=Workbook()
HF=PatternFill('solid',fgColor='1F3A5F'); HFont=Font(bold=True,color='FFFFFF')
def header(ws,cols,widths):
    ws.append(cols)
    for i,c in enumerate(cols,1):
        cell=ws.cell(1,i); cell.fill=HF; cell.font=HFont; cell.alignment=Alignment(vertical='center',wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1] if i-1<len(widths) else 18
    ws.freeze_panes='A2'
ws=wb.active; ws.title='Contactos'
COLS=['ID','Nombre','Cargo','Empresa','País','Email','LinkedIn','Origen','Estado','Prioridad','Próxima acción','Notas','Rubro(s)','Segmento SafetyMind','Teléfono(s)','Web','Dirección','Ciudad/Región','Anunciante']
header(ws,COLS,[6,22,20,40,10,32,14,20,12,10,45,60,45,38,30,30,50,14,11])
for i,r in enumerate(rows,1):
    ws.append([i]+[r[c] for c in COLS[1:]])
    ws.cell(ws.max_row,15).number_format='@'
n=len(rows)+1
t=Table(displayName='Contactos',ref=f"A1:{get_column_letter(len(COLS))}{n}"); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
dv1=DataValidation(type='list',formula1='"Prospecto,Contactado,Reunión,Propuesta,Cliente,No interesado"',allow_blank=True); dv1.add(f'I2:I{n+500}')
dv2=DataValidation(type='list',formula1='"Alta,Media,Baja,Por calificar"',allow_blank=True); dv2.add(f'J2:J{n+500}')
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
ws2=wb.create_sheet('Oportunidades')
header(ws2,['ID oportunidad','ID contacto','Empresa','Necesidad','Solución propuesta','Etapa','Valor estimado','Probabilidad','Próxima acción','Fecha'],[14,12,40,40,40,14,14,12,40,12])
ws3=wb.create_sheet('Resumen')
header(ws3,['Evento/Fuente','País','Cantidad','',''],[50,16,12,12,12])
paises=['Perú','Colombia','Brasil','Argentina','Canadá','Estados Unidos','Australia','China']
for i,p in enumerate(paises,2): ws3.append(['Directorio SSO 2026',p,f'=COUNTIFS(Contactos!$H:$H,$A{i},Contactos!$E:$E,$B{i})'])
ws3.append(['Total','',f'=SUM(C2:C{len(paises)+1})']); ws3.append([''])
def hdr(row,n):
    for c in range(1,n+1): ws3.cell(row,c).fill=HF; ws3.cell(row,c).font=HFont
r0=ws3.max_row+1; ws3.append(['Segmento SafetyMind','Alta','Media','Baja','Total']); hdr(r0,5)
for s_ in ORDER:
    i=ws3.max_row+1
    ws3.append([s_]+[f'=COUNTIFS(Contactos!$N:$N,$A{i},Contactos!$J:$J,{col}${r0})' for col in 'BCD']+[f'=SUM(B{i}:D{i})'])
i=ws3.max_row+1; ws3.append(['Total']+[f'=SUM({c}{r0+1}:{c}{i-1})' for c in 'BCDE']); ws3.append([''])
r1=ws3.max_row+1; ws3.append(['Estado (pipeline)','Cantidad']); hdr(r1,2)
for e_ in ['Prospecto','Contactado','Reunión','Propuesta','Cliente','No interesado']:
    i=ws3.max_row+1; ws3.append([e_,f'=COUNTIF(Contactos!$I:$I,A{i})'])
ws3.append(['']); r2=ws3.max_row+1; ws3.append(['Indicadores de calidad de datos','Cantidad']); hdr(r2,2)
for lab,f in [('Con email','=COUNTIF(Contactos!$F$2:$F$2000,"?*")'),('Con teléfono','=COUNTA(Contactos!$O$2:$O$2000)'),('Con contacto nominal','=COUNTIF(Contactos!$B$2:$B$2000,"?*")'),('Anunciantes del directorio','=COUNTIF(Contactos!$S:$S,"Sí")'),('Con LinkedIn validado','=COUNTIF(Contactos!$G$2:$G$2000,"?*")')]:
    ws3.append([lab,f])
ws4=wb.create_sheet('Eventos')
header(ws4,['Evento','País','Valor para SafetyMind'],[50,12,90])
ws4.append(['Directorio SSO 2026 – Cero Accidentes (22.ª ed.)','Perú','Mapa de ~670 proveedores SST en Perú: contratistas mineros (clientes potenciales), consultoras SST (canal), TI/telecom (partners/competencia), distribuidores EPP (canal hacia minería) y 21 anunciantes con inversión activa en SST.'])
wb.save(sys.argv[2])
c=collections.Counter(r['Prioridad'] for r in rows); print(len(rows),c, sum(1 for r in rows if r['Nombre']), sum(1 for r in rows if r['Email']))
print(collections.Counter(r['Segmento SafetyMind'] for r in rows))
