import csv, re, sys, json, collections, unicodedata
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
# usage: build_exponor.py exponor.csv CRM_Directorio_SSO_2026.xlsx CRM_ExpoMina_2026.xlsx out.xlsx
src, dirxlsx, expoxlsx, out = sys.argv[1:5]
def clean(s): return re.sub(r'\s+',' ',s.replace('\xa0',' ')).strip()
def norm(s): return unicodedata.normalize('NFKD',s.upper()).encode('ascii','ignore').decode()
def key(n):
    n=norm(n).lower()
    n=re.sub(r'\b(s\.?\s?a\.?\s?c\.?|s\.?\s?a\.?|e\.?i\.?r\.?l\.?|s\.?r\.?l\.?|ltda\.?|limitada|spa|sac|eirl|srl|peru|chile|del|de|y|cia|co|ltd|inc|corp|products)\b','',n)
    return re.sub(r'[^a-z0-9]','',n)
def load(path,keycol='Empresa'):
    ws=openpyxl.load_workbook(path)['Contactos']; h=[c.value for c in ws[1]]; m={}
    for r in ws.iter_rows(min_row=2,values_only=True):
        d=dict(zip(h,r))
        if d.get(keycol): m.setdefault(key(d[keycol]),d)
    return m
dirmap=load(dirxlsx); expomap=load(expoxlsx)

TAGC={'China':'China','Alemania':'Alemania','Estados Unidos':'Estados Unidos','Perú':'Perú','Canadá':'Canadá','Provincia de Ontario':'Canadá',
 'Australia':'Australia','Brasil':'Brasil','Italia':'Italia','Francia':'Francia','España':'España','Suiza':'Suiza','Países Bajos':'Países Bajos',
 'Argentina':'Argentina','Bélgica':'Bélgica','Finlandia':'Finlandia','Suecia':'Suecia','Polonia':'Polonia','Reino Unido':'Reino Unido',
 'República Checa':'República Checa','Hungría':'Hungría'}
rows=list(csv.DictReader(open(src,encoding='utf-8-sig')))
comp=collections.OrderedDict()
for r in rows:
    raw=clean(r['Empresa'])
    ZONES=set(TAGC)|{'EstaremosEnExponor','EstaremosenExponor','Lanza tu Innovación','AIA-Pyme','de Energía','Energía'}
    def known(t): return clean(t).replace('Pabellón ','').replace('Pabellon ','') in ZONES
    tags=[t for t in re.findall(r'\(\s*([^()]*)\)',raw) if known(t)]
    name=clean(re.sub(r'\(\s*([^()]*)\)',lambda m:'' if known(m.group(1)) else m.group(0),raw)).rstrip(' .') or raw
    if re.search(r'Sieyuan',raw): name='Sieyuan Electric Co., Ltd.'
    k=key(name)
    c=comp.setdefault(k,{'Empresa':name,'stands':[],'tags':[]})
    pab=clean(r['Pabellón']).replace('Pabellon','Pabellón').replace('Pabellón ','')
    st=f"{pab}-{clean(r['ID_Pabellón'])}"
    if st not in c['stands']: c['stands'].append(st)
    for t in tags:
        t=clean(t).replace('Pabellón ','').replace('Pabellon ','')
        t={'EstaremosenExponor':'EstaremosEnExponor','de Energía':'Energía'}.get(t,t)
        if t and t not in c['tags']: c['tags'].append(t)

CHN=r'SHANDONG|ZHEJIANG|JIANGSU|HEBEI|HENAN|ZHENGZHOU|SHANGHAI|BEIJING|GUANGDONG|GUANGZHOU|SHENYANG|TIANJIN|QINGDAO|WUXI|CHANGSHA|JIAOZUO|BAOJI|LUOYANG|HUNAN|HUBEI|ANHUI|FUJIAN|SHAANXI|SHANXI|NINGBO|NINGGUO|YANTAI|JINAN|XUZHOU|SICHUAN|LIAONING|DALIAN|JIANGXI|SHENZHEN|DONGGUAN|HEFEI|ZHUZHOU|TAIZHOU|WENZHOU|WEIHAI|ANSHAN|DANDONG|HAICHENG|HANDAN|HUAINAN|KUNSHAN|YINGKOU|YUNCHENG|ZHUCHENG|JINGJIANG|QUANZHOU|SHIJIAZHUANG|YANGGU|HONGKONG'
def pais(name,tags):
    for t in tags:
        for kk,v in TAGC.items():
            if t.startswith(kk) or t==kk: return v
    n=norm(name)
    if re.search(PERU,n): return 'Perú'
    if re.search(CHN,n): return 'China'
    if 'CHILE' in n or re.search(r'\bSPA\b|\bLTDA\b|LIMITADA|\bSCM\b|ANTOFAGASTA',n): return 'Chile'
    if 'GMBH' in n: return 'Alemania'
    if re.search(r'SP\. Z O\.O',n): return 'Polonia'
    if re.search(r'\bE\.?I\.?R\.?L\b|\bS\.?A\.?C\b',n): return 'Perú (por validar)'
    if re.search(r'\bINC\b|\bLLC\b|\bCORP\b',n): return 'Estados Unidos (por validar)'
    if re.search(r'CO\.?,? ?LTD',n): return 'China (por validar)'
    return 'Chile'

INST=r'ASIA CONSULTING|EMBAJADA|EMBASSY|CAMARA|CHAMBER|PROCHILE|PROMPERU|PROMARGENTINA|\bICEX\b|AGENCIA ITALIANA|\bAWEX\b|AUSTMINE|ADVANTAGE AUSTRIA|\bAHK\b|JAPAN EXTERNAL|TRADE AND INVESTMENT|INVEST QUEBEC|OPPORTUNITIES NEW BRUNSWICK|GOBIERNO|MUNICIPALIDAD|DELEGACION PRESIDENCIAL|REGION DE LOS|UNION EUROPEA|CLUSTER|PROLOA|ASOCIACION|CONSEJO|ABIMAQ|ADIMRA|INSTITUTO DE INGENIEROS|APRIMIN|^AIE$|EXPOMIN 20|POLSKA AGENCJA|FUERZA AEREA|COLEGIO|LICEO|MSTA CANADA|NETHERLANDS LOUNGE|MESSE|FUNDACION|MERCURIO|REVISTA|RADIO ANTOFAGASTA|REPORTE MINERO|MEDIO MINERIA|EDITORIAL|MULTIMEDIOS|MINEGROUP CONGRESOS|MEETING POINT|PUEBLOS ATACAMENOS'
PERU=r'INGENIEROS DE MINAS DEL PERU'
MINE=r'CODELCO|^BHP$|ANTOFAGASTA MINERALS|^SQM$|ALBEMARLE|^GLENCORE$|\bTECK\b|SIERRA GORDA|CAPSTONE|MINERA EL ABRA|COMPANIA MINERA|CIA\. MINERA|NOVANDINA|CERRO ALTO MINERIA'
EPP=r'^3M$|^MSA\b|DRAGER|VICSA|BOLLE SAFETY|DELTA PLUS|IRON SAFETY|HD SAFETY|BOYUAN SAFETY|BATA INDUSTRIAL|WORKING APPAREL|CONFECCIONES|DONNELY|TEXMUNDO|MARITEX|ALCOSAFE|JAYSON|\bEPP\b|PROTECCION PERSONAL|INDUMENTARIA|UNIFORM'
SEG=r'SEGURIDAD|FIRE|INCENDIO|PROSEGUR|NORSEG|SEGUSA|TECNISEC|DIPROSEG|SEGURYCEL|SAFE-T|SAFE GAUGE|SAFESMART|MINEARC|GREEN & SAFE|ANTICOLISION|FATIG|SAFETREE|RESCATE|SENALIZ|SAFEM\b|SAFEME|GAS DETECT|LOCKOUT|BLOQUEO'
CONTR=r'STRACON|THIESS|SIGDO KOPPERS|ECHEVERRIA IZQUIERDO|ICIL ICAFAL|CONFIPETROL|MAMMOET|SACYR|SOLETANCHE|ENAEX|ORICA|DYNO NOBEL|FAMESA|ROCKBLAST|EXPLOSIV|CONSTRUC|MONTAJE|MANTENCION|MANTENIMIENTO|SERVICIOS MINEROS|MINING SERVICES|MINING SUPPORT|TRANSPORTE|BUSES HUALPEN|MINABUS|BUS 4X4|SOTRASER|RODOTRANS|SITRANS|GRUAS|CRANES|IZAJE|RENTAL|ANDAMIOS|LAYHER|^FCAB$|FERRONOR|PUERTO DE|PORTUARIO|KALTIRE|GRUPO EULEN|SERVIALL|VERTICAL MINE|CHRISTENSEN|INDEPENDENT DRILLING|SWASTIC DRILLING|ALL DRILL|LOGISTIC|SOUTH CRANES|SMARTLIFT|TODO IZAJE|CEMTEC|BCM SERVICIOS|CAS SERVICIOS INDUSTRIALES|INPPA SERVICIOS|DYC SERVICIOS|MOVERENT|VERTEX RENTAL|DEMOLITION|RMI CONSTRUC|UNIDAD CONSTRUC|DURHAM|M4TS'
OEM=r'KOMATSU|EPIROC|SANDVIK|CATERPILLAR|FINNING|LIEBHERR|XCMG|^SANY|METSO|FLSMIDTH|\bWEIR\b|ATLAS COPCO|BOART LONGYEAR|MICHELIN|^JCB$|MAN Y VOLKSWAGEN|^TOYOTA$|MITSUB|CUMMINS|TAKRAF|TOMRA|HERRENKNECHT|GOLDHOFER|FAYMONVILLE|FORKLIFT|^CAEX|^DETROIT$|LINGLONG|MI-JACK|COMBILIFT|NORMET|HITACHI|VOLVO|SCANIA|WABTEC|JIKAI|SUNWARD|RAILWAY CONSTRUCTION HEAVY'
TECH=r'GEOTAB|WEBFLEET|SITRACK|WISETRACK|HEXAGON|LEICA|MAPTEK|DATAMINE|MOTIONMETRICS|MOVISIGHT|PITCREW|NEOS AI|NEURO4B|NEURITIVA|INTELLEGERE|TIMINING|FRACTTAL|^SONDA$|^ENTEL$|GRUPO GTD|AXIS COMMUNICA|HANWHA VISION|CONVERGINT|SMART COMPLIANCE|KRUX ANALYTICS|GEOLOGICAI|SPARTANLYNC|HIVERADAR|SKKYNET|YOY SIMULATORS|SINGULARISVR|ROBOTIA|RUKIAN|ZUVAR|SENTRYAPP|ISANTO|IAP SOLUTIONS|DATASCOPE|AUTOMINING|FAST2 MINE|TECWISE|VERIDAPT|TIVIT|DEVCODERX|BQSENSE|GUNJOP|HI MINING|^TOMOS$|^SICK|YOKOGAWA|DJI|ATACAMACOPTER|UNMANNED AERIAL|WINGTRA|SKYGEO|DEEP TREKKER|MINESENSE|ZIGOMIN|AMPERXGROUP|SOFTWARE|DIGITAL|ANALYTIC|\bIOT\b|\bAI\b|INTELIGEN|INTELLIGENCE|TELEMETR|MONITOREO|MONITORING|TELECOMUNICACION|COMUNICACION|RADIOCOMUNIC|SATELIT|\bDRON|ROBOT|VISION|SIMULATOR|SENSOR|CONECTIV|NETTZ|TELTRONIC|ELECTRORAM|MUNDO TELECOM|MKS SOLUCIONES|KIZEO|BUK$|TSCOM|INNERVYCS|AMV LATAM|WARRIORS INNOVA|PULSO SPA|ZIZCAR|MINERIA XR|CLOUDCSP|INFOTECH|TRIVICK|FASTECH|INBOX TECHNOLOGY|HELPTECH|BEE3 MINING|LAVOR MINING'
CONS=r'CONSULT|CAPACITA|\bOTEC\b|ASESOR|CERTIFIC|BUREAU VERITAS|ALS INSPECTION|DICTUC|IDIEM|UNIVERSIDAD|ACADEMY|LEAN INSTITUTE|CEFOMIN|ROCKTRAIN|XP ONBOARDING|M-RISK|HAZARD CONTROL|^INERCO$|^RINA|ARCADIS|^BBA$|CENTRO DE INNOVACION|ADVANCED MINING TECHNOLOGY|ISOTEMPO|SOLUTIVA|INGENIERIA EN EVOLUCION|FULCRO|CENTRO CHILEVALORA|OTIC|PETRA ACADEMY|INSPECC|AUDIT|COPPER CAPACITACION|CAPELEC|SEGUROS|CORREDORES'
SAL=r'SALUD|MEDIC|CLINICA|PREOMED|CARDIO|BRAIN TREATMENT|BUPA|MUTUAL|\bACHS\b|LABORAL|SALUDABLE'
SEGS=[('Institución / gremio / pabellón país',INST,'Media','Networking institucional; identificar contacto de pabellón'),
 ('Minera (cliente final)',MINE,'Alta','Mapear Gerente SSOMA/HSE y faenas; proponer piloto'),
 ('Canal / distribuidor EPP',EPP,'Media','Evaluar como canal de distribución hacia minería/industria'),
 ('Aliado complementario – seguridad',SEG,'Media','Evaluar co-venta o integración de soluciones'),
 ('Cliente potencial – contratista / operación de alto riesgo',CONTR,'Alta','Identificar Gerente HSE/SSOMA en LinkedIn y proponer demo'),
 ('Partner OEM – equipos mineros',OEM,'Media','Explorar integración con flota/equipos (visión + telemetría)'),
 ('Partner tecnológico / competidor',TECH,'Alta','Mapear oferta: evaluar integración o competencia'),
 ('Canal / aliado – consultoría SST',CONS,'Alta','Explorar alianza comercial / reventa a sus clientes'),
 ('Aliado complementario – salud',SAL,'Media','Evaluar alianza en salud ocupacional / fatiga')]
ORDER=['Minera (cliente final)','Cliente potencial – contratista / operación de alto riesgo','Partner tecnológico / competidor','Canal / aliado – consultoría SST',
 'Partner OEM – equipos mineros','Canal / distribuidor EPP','Aliado complementario – seguridad','Aliado complementario – salud','Institución / gremio / pabellón país','Proveedor industrial (equipos/insumos)']
ACC={s[0]:s[3] for s in SEGS}; ACC['Proveedor industrial (equipos/insumos)']='Nutrir con contenido; sin acción comercial inmediata'
PRI={s[0]:s[2] for s in SEGS}; PRI['Proveedor industrial (equipos/insumos)']='Baja'
DSEG_MAP={'Cliente potencial – operación de alto riesgo':'Cliente potencial – contratista / operación de alto riesgo'}
out_rows=[]
for k,c in comp.items():
    n=norm(c['Empresa']); p=pais(c['Empresa'],c['tags'])
    d=dirmap.get(k); e=expomap.get(k)
    seg=None
    for s,pat,_,_ in SEGS:
        if re.search(pat,n):
            if p.startswith('China') and s in ('Cliente potencial – contratista / operación de alto riesgo','Minera (cliente final)','Aliado complementario – seguridad'): continue
            seg=s; break
    if p.startswith('China') and seg=='Partner tecnológico / competidor' and not re.search(r'DIGITAL|SOFTWARE|\bAI\b|VISION|IOT|ROBOT|INTELLIGEN',n): seg=None
    if seg is None and e and e['Segmento SafetyMind'] in ORDER and e['Segmento SafetyMind']!='Proveedor industrial (equipos/insumos)': seg=e['Segmento SafetyMind']
    if seg is None and d:
        ds=DSEG_MAP.get(d['Segmento SafetyMind'],d['Segmento SafetyMind'])
        if ds in ORDER: seg=ds
    if seg is None: seg='Proveedor industrial (equipos/insumos)'
    pri=PRI[seg]
    src=d or {}
    notas="Exponor 2026."
    if len(c['stands'])>1: notas+=" Varios stands (mayor inversión)."
    tz=[t for t in c['tags'] if t in ('Lanza tu Innovación','EstaremosEnExponor','EstaremosenExponor','AIA-Pyme','de Energía','Energía')]
    if 'Lanza tu Innovación' in tz: notas+=" Startup (zona Lanza tu Innovación)."
    if 'AIA-Pyme' in tz: notas+=" Pyme regional (pabellón AIA)."
    if any('Energía' in t for t in tz): notas+=" Pabellón de Energía."
    cruce=[]
    if d: cruce.append(f"Directorio SSO (ID {d['ID']})"); notas+=f" También en CRM Directorio SSO 2026 (ID {d['ID']})"+("; contacto heredado = filial Perú." if not p.startswith('Perú') else ".")
    if e: cruce.append(f"ExpoMina (ID {e['ID']})"); notas+=f" También expone en ExpoMina 2026 (ID {e['ID']})."
    if '(por validar)' in p: notas+=" País inferido; validar."
    if not d: notas+=" Sin contacto publicado."
    out_rows.append({'Nombre':src.get('Nombre') or '','Cargo':src.get('Cargo') or '','Empresa':c['Empresa'],'País':p,
      'Email':src.get('Email') or '','LinkedIn':'','Origen':'Exponor 2026','Estado':'Prospecto','Prioridad':pri,
      'Próxima acción':ACC[seg],'Notas':notas,'Segmento SafetyMind':seg,'Stand(s)':', '.join(c['stands']),'N° stands':len(c['stands']),
      'Pabellón país / zona':', '.join(c['tags']),'Web':src.get('Web') or (e or {}).get('Web') or '','Teléfono(s)':src.get('Teléfono(s)') or '',
      'Cruce otros CRM':'; '.join(cruce) or 'No'})
porder={'Alta':0,'Media':1,'Baja':2}
out_rows.sort(key=lambda r:(porder[r['Prioridad']],ORDER.index(r['Segmento SafetyMind']),-r['N° stands'],norm(r['Empresa'])))
COLS=['ID','Nombre','Cargo','Empresa','País','Email','LinkedIn','Origen','Estado','Prioridad','Próxima acción','Notas','Segmento SafetyMind','Stand(s)','N° stands','Pabellón país / zona','Web','Teléfono(s)','Cruce otros CRM']
wb=Workbook(); HF=PatternFill('solid',fgColor='1F3A5F'); HFont=Font(bold=True,color='FFFFFF')
def header(ws,cols,widths):
    ws.append(cols)
    for i,cn in enumerate(cols,1):
        cell=ws.cell(1,i); cell.fill=HF; cell.font=HFont; cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1] if i-1<len(widths) else 16
    ws.freeze_panes='A2'
ws=wb.active; ws.title='Contactos'
header(ws,COLS,[6,20,20,45,18,30,14,16,12,10,45,55,40,22,9,24,32,28,28])
for i,r in enumerate(out_rows,1):
    ws.append([i]+[r[c] for c in COLS[1:]]); ws.cell(ws.max_row,18).number_format='@'
n=len(out_rows)+1
t=Table(displayName='Contactos',ref=f"A1:{get_column_letter(len(COLS))}{n}"); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
dv1=DataValidation(type='list',formula1='"Prospecto,Contactado,Reunión,Propuesta,Cliente,No interesado"',allow_blank=True); dv1.add(f'I2:I{n+500}')
dv2=DataValidation(type='list',formula1='"Alta,Media,Baja,Por calificar"',allow_blank=True); dv2.add(f'J2:J{n+500}')
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
ws2=wb.create_sheet('Oportunidades'); header(ws2,['ID oportunidad','ID contacto','Empresa','Necesidad','Solución propuesta','Etapa','Valor estimado','Probabilidad','Próxima acción','Fecha'],[14,12,40,40,40,14,14,12,40,12])
ws3=wb.create_sheet('Resumen'); header(ws3,['Evento/Fuente','País','Cantidad','',''],[50,26,12,12,12])
pc=collections.Counter(r['País'] for r in out_rows)
for i,(p,_) in enumerate(pc.most_common(),2): ws3.append(['Exponor 2026',p,f'=COUNTIFS(Contactos!$H:$H,$A{i},Contactos!$E:$E,$B{i})'])
last=ws3.max_row; ws3.append(['Total','',f'=SUM(C2:C{last})']); ws3.append([''])
def hdr(row,k):
    for cc in range(1,k+1): ws3.cell(row,cc).fill=HF; ws3.cell(row,cc).font=HFont
r0=ws3.max_row+1; ws3.append(['Segmento SafetyMind','Alta','Media','Baja','Total']); hdr(r0,5)
for s in ORDER:
    i=ws3.max_row+1; ws3.append([s]+[f'=COUNTIFS(Contactos!$M:$M,$A{i},Contactos!$J:$J,{col}${r0})' for col in 'BCD']+[f'=SUM(B{i}:D{i})'])
i=ws3.max_row+1; ws3.append(['Total']+[f'=SUM({cl}{r0+1}:{cl}{i-1})' for cl in 'BCDE']); ws3.append([''])
r1=ws3.max_row+1; ws3.append(['Estado (pipeline)','Cantidad']); hdr(r1,2)
for e_ in ['Prospecto','Contactado','Reunión','Propuesta','Cliente','No interesado']:
    i=ws3.max_row+1; ws3.append([e_,f'=COUNTIF(Contactos!$I:$I,A{i})'])
ws3.append(['']); r2=ws3.max_row+1; ws3.append(['Indicadores','Cantidad']); hdr(r2,2)
for lab,f in [('Registros originales (stands)',len(rows)),('Empresas únicas','=COUNTA(Contactos!$D$2:$D$3000)'),('También en ExpoMina 2026','=COUNTIF(Contactos!$S$2:$S$3000,"*ExpoMina*")'),('También en Directorio SSO','=COUNTIF(Contactos!$S$2:$S$3000,"*Directorio*")'),('Startups (Lanza tu Innovación)','=COUNTIF(Contactos!$P$2:$P$3000,"*Lanza tu Innovación*")'),('Con 2+ stands','=COUNTIF(Contactos!$O$2:$O$3000,">1")')]:
    ws3.append([lab,f])
ws4=wb.create_sheet('Eventos'); header(ws4,['Evento','País','Valor para SafetyMind'],[45,12,100])
ws4.append(['Directorio SSO 2026 – Cero Accidentes (22.ª ed.)','Perú','Mapa de ~670 proveedores SST en Perú (ver CRM_Directorio_SSO_2026).'])
ws4.append(['ExpoMina Perú 2026','Perú','Feria minera: mineras, contratistas, OEM, tecnología y proveedores (ver CRM_ExpoMina_2026).'])
ws4.append(['Exponor 2026 (Antofagasta)','Chile','Principal feria minera del norte de Chile: mineras (Codelco, BHP, AMSA, SQM, Albemarle), contratistas, OEM, startups y pabellones país; base para reuniones en stand y prospección post-evento.'])
wb.save(out)
print(len(rows),len(out_rows)); print(collections.Counter(r['Segmento SafetyMind'] for r in out_rows)); print(collections.Counter(r['Prioridad'] for r in out_rows)); print(pc.most_common(30))
print('dir',sum('Directorio' in r['Cruce otros CRM'] for r in out_rows),'expo',sum('ExpoMina' in r['Cruce otros CRM'] for r in out_rows))
json.dump(out_rows,open(out+'.json','w'),ensure_ascii=False)
