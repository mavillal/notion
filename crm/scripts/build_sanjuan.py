import csv, re, sys, json, collections, unicodedata
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
# usage: build_sanjuan.py sj.csv Directorio.xlsx ExpoMina.xlsx Exponor.xlsx AquaSur.xlsx FIDAE.xlsx out.xlsx
src, out = sys.argv[1], sys.argv[-1]
def clean(s): return re.sub(r'\s+',' ',(s or '').replace('\xa0',' ')).strip()
def norm(s): return unicodedata.normalize('NFKD',s.upper()).encode('ascii','ignore').decode()
def key(n):
    n=norm(n).lower()
    n=re.sub(r'\(.*?\)','',n)
    n=re.sub(r'\b(s\.?\s?a\.?\s?c\.?|s\.?\s?a\.?\s?u\.?|s\.?\s?a\.?|s\.?\s?a\.?\s?s\.?|e\.?i\.?r\.?l\.?|s\.?r\.?l\.?|s\.?r\.?o\.?|s\.?l\.?|ltda\.?|limitada|spa|sac|sas|saic|saci|eirl|srl|peru|chile|argentina|del|de|y|cia|co|ltd|inc|corp|corporation|company|group|grupo|gmbh|ag|as|llc|plc)\b','',n)
    return re.sub(r'[^a-z0-9]','',n)
def load(path):
    ws=openpyxl.load_workbook(path)['Contactos']; h=[c.value for c in ws[1]]; m={}
    for r in ws.iter_rows(min_row=2,values_only=True):
        d=dict(zip(h,r))
        if d.get('Empresa') and key(d['Empresa']): m.setdefault(key(d['Empresa']),d)
    return m
NAMES=['Directorio SSO','ExpoMina','Exponor','AquaSur','FIDAE']
CRMS=[(n,load(p)) for n,p in zip(NAMES,sys.argv[2:7])]

DELEG={'DELEGACION PERU':'Perú','DELEGACION CANADA':'Canadá','NETHERLANDS':'Países Bajos','DELEGACION ALEMANA':'Alemania'}
PAIS_FIX={'FUGRO CHILE S.A.':'Chile','HERRENKNECHT CHILE LTDA.':'Chile','RELIPER':'Chile','ZIGMA SOLUCIONES INDUSTRIALES LTDA':'Chile',
 'CASTEM':'Perú','FRESOMAC SAC':'Perú','JSIROCK TOOLS CO.,LTD':'China','QINGDAO FAMBITION':'China','GLOBALBRING ASESORIA':'China',
 'SABLE RESOURCES LTD':'Canadá','FINLANDIA':'Finlandia','MAGGIORA LLC':'Estados Unidos','A.G. DER DILLINGER HÜTTENWERKE':'Alemania',
 'DELEGACIÓN ALEMANA - MESSE DUESSELDORF':'Alemania','DELEGACIÓN CANADÁ':'Canadá','DELEGACION PERU':'Perú','NETHERLANDS / PAISES BAJOS':'Países Bajos',
 'BOSCH REXROTH S.A.I.C.':'Argentina','IFM ELECTRONIC SRL':'Argentina','PROMINENT ARGENTINA S.A.':'Argentina','WILO SALMSON ARGENTINA S. A.':'Argentina',
 'ONTEC FORTINOX S.A.U.':'Argentina','PERI S.A.':'Argentina','RANDSTAD':'Argentina','DHL GLOBAL FORWARDING':'Argentina','STEINERT LATINOAMERICANA LTDA.':'Chile'}

S_MINE='Cliente final – minera / energía (operación o proyecto)'
S_CONTR='Cliente final – contratista minero (operación, perforación, transporte, campamento)'
S_TECH='Partner tecnológico / competidor'
S_CONS='Canal / aliado – consultoría, capacitación y certificación'
S_OEM='Partner OEM – maquinaria y equipos mineros'
S_EPP='Canal / distribuidor EPP'
S_SEG='Aliado complementario – seguridad y emergencias'
S_SAL='Aliado complementario – salud'
S_INST='Institución / cámara / medio / gobierno'
S_PROV='Proveedor industrial / servicios'
S_OUT='Fuera de foco (turismo y servicios locales)'
CUR={
 S_MINE:['ALDEBARAN RESOURCES INC.','CALERAS SAN JUAN','CALIDRA','CASPOSO ARGENTINA LTD. - SUCURSAL ARGENTINA','GLENCORE EL PACHÓN','GOLDEN MINING',
  'MCEWEN COPPER PROYECTO LOS AZULES','MOGOTES METALS INC.','NEW PEAK','SABLE RESOURCES LTD','VELADERO','VICUÑA','YPF','PAN AMERICAN ENERGY - AXION PUEBLA','EPSE','ECO MINERA'],
 S_CONTR:['ACONCAGUA SERVICIOS MINEROS SRL','ADL','ANDEMIRA S.A.','C&M SA - CONSTRUCCION Y MINERIA SA','CONOSUR DRILLING S.A.','DRILLAR SAN JUAN','ENAEX','ORICA',
  'MAJOR DRILLING','HG PERFORACIONES','DELTA MINING','ECOVIAL','MILICIC S.A.','MAPAL SIGMA - MINING CONSTRUCTIONS','MINEXO','ZLATO/DUMANDZIC','CINTER','INTERENERGY',
  'BECHER LOGISTICA','SOTUR','TRANSPORTES PAP','TECMACO','GRÚAS SAN BLAS','GRUAS BILLIA','GRUAS MARTIN','ARAMARK','CATERWEST','ATA','JSC SOLUCIONES INTEGRALES',
  'NIELSEN EXPEDICIONES','TAGING','SYOTEC','GSI','ANDESITA SAS','CASTEM','SERVICIOS MINEROS DE AMERICA','TERRA LOGISTICA','OAPCE MULTITRANS Y CARGO TRANSFER CUYANA',
  'L & G LOGISTICA Y GASTRONOMIA','NAV CAMPAMENTOS','COCYAR S.A.','WINDLAND','ESPECO','CRUZ DEL SUR','TRANSPORTE VESPRINI','FG LOGISTICA','PUERTOTRANS','GEOTUB',
  'JRI ARGENTINA','INGEAP','GRUPO SIT SERVICIOS INTEGRALES TOPOGRÁFICOS','B & W','ARCE GEOFÍSICOS ARGENTINA','GEOAR S.R.L','PARAMASSI ANDINA SRL','CUYO MINERO'],
 S_TECH:['TCV DESARROLLO DE SOLUCIONES','MAKERS SAS','SITRACK','GEOTAB','MINETECH','MC MONITOREO & CONTROL DE PERFORACION','GHM SATELITAL SRL','TESACOM','TCTECH SRL',
  'TECHNOLOGY BUREAU - NOKIA','MARIO ZINI & CIA S.R.L','INTERSAT SA','DATAWISE','DEKA LYSEIS','TRIM STUDIOS SRL - SYTEX','ORIGAMI SOFTWARE','I-POINT S. A. S.',
  'GEOSISTEMAS','ALS SOLUCIONES TECNOLOGICAS','GAMT','ISTRA','SLIMSTOCK (ALAMO)','DISEGNO','GREENWORKING'],
 S_CONS:['TÜV RHEINLAND ARGENTINA','CONSULTORA OLGUÍN','MAGISTRALI CONSULTORES','MPC CONSULTORA','AMERICAN ADVISOR','RANDSTAD','MANPOWER','UNIVERSIDAD DE SAN JUAN',
  'ARCADIS','HIDROAR S.A.','ENVIRO S G','INGENIERIA INTEGRAL'],
 S_OEM:['EPIROC','SANDVIK','KOMATSU ISEMAR','FINNING','METSO ARGENTINA SA','FLSMIDTH.','WEIR','ME ELECMETAL','BRADKEN','RESEMIN','VOLVO','ESCANDINAVIA DEL PLATA',
  'LOPEZ ARGENTINA','SCANIA','VOLKSWAGEN ARGENTINA S.A.','FICAMEN','DIESEL LANGE','COVEMA','BAUZA','PALFINGER COMPANY','SULLAIR ARGENTINA','KAESER COMPRESORES DE ARGENTINA SA',
  'LAROCCA MINERÍA','H&B','WESTPRO MACHINERY INC.','FOREMOST INDUSTRIES LP','HERRENKNECHT CHILE LTDA.','JSIROCK TOOLS CO.,LTD','QINGDAO FAMBITION','GLOBALBRING ASESORIA',
  'DISTRIZAN','BOSCH REXROTH S.A.I.C.','SIEMENS','STEINERT LATINOAMERICANA LTDA.','TOMRA SORTING GMBH','ENDUROCO','DOBLE TRACCION','RUBIA AUTOMOTORES SRL','MELEX ARGENTINA'],
 S_EPP:['A. MARSHALL MOFFAT','ARRAIGO','BORIS HERMANOS','RINGO WORK','YARD','SOLIMIN','MAINCAL SA'],
 S_SEG:['ECA FIRE/ALR SYSTEMS SISTEMAS CONTRA INCENDIOS','GRUPO OMEGA - S3 S.R.L.','GLOBAL MINING - MAQUINARIAS RAWSON','HUARPE SEGURIDAD INTEGRAL','INSTALROS',
  'PRODIMIN','TECIN MINERÍA SRL','RIVADAVIA SEGUROS'],
 S_SAL:['CLINICA ALEMANA','SUMAR SALUD','EMERGENCIAS'],
 S_INST:['AHK - CÁMARA DE INDUSTRIA Y COMERCIO ARGENTINO-ALEMANA','ARMINERA','ASOCIACIÓN OBRERA MINERA ARGENTINA','CAMARA ARGENTINA DE LA CONSTRUCCIÓN',
  'CÁMARA MINERA DE SAN JUAN','CAMARA PROVEEDORES MINEROS DE USPALLATA','CAPRIMSA','CASEMI','CASETIC','CFI - CONSEJO FEDERAL DE INVERSIONES','CLUSTER BERAZATEGUI',
  'CLÚSTER DE PETRÓLEO GAS Y MINERÍA DE CORDOBA','DELEGACIÓN ALEMANA - MESSE DUESSELDORF','DELEGACIÓN CANADÁ','DELEGACION PERU','FINLANDIA','GAPP','GOBIERNO DE SAN JUAN',
  'MINISTERIO DE MINERIA - GOBIERNO DE CATAMARCA','NETHERLANDS / PAISES BAJOS','PROMARGENTINA','UNIÓN INDUSTRIAL SAN JUAN','PANORAMA MINERO SA','PETROQUÍMICA',
  'REVISTA ENERGIA EN MOVIMIENTO','GSMEDIA','CADENA3','CANAL 13 Y DIARIO 13 SAN JUAN','CANAL 8','TELESOL-ZONDA','RADIO COM','BATALLER CONTENIDOS','FEDERACION ECONOMICA SAN JUAN',
  'CAMARA COMERCIANTES UNIDOS SAN JUAN','BANCO SAN JUAN S.A.','BANCO SANTANDER','BANCO SUPERVIELLE','ALLARIA S.A'],
 S_OUT:['CAMARA DE PRESTADORES TURISTICOS DE LA PROVINCIA DE SAN JUAN','CHAMPAÑERA MIGUEL MAS','ESCAPE AVENTURA','FINCA DON ELIAS','FLORERIA ARGENTINA',
  'GARDENIA ALQUILER TEMPORARIO','IGARRETA SACI','INTERLAGOS ADV','LA SOÑADA','MUNICIPALIDAD DE VALLE FERTIL','PARQUE PARADOR','YAFAR DESTINOS','GREENFIELD MINING',
  'RE/MAX CUYANA','AUTOPARK','EL BOLSILLO','FG PRODUCTORA','WE MADE','CASABLANCA','BRANDS','LEPLAG FUMIGACIONES'],
}
ORDER=[S_MINE,S_CONTR,S_TECH,S_CONS,S_OEM,S_EPP,S_SEG,S_SAL,S_INST,S_PROV,S_OUT]
PRI={S_MINE:'Alta',S_CONTR:'Alta',S_TECH:'Alta',S_CONS:'Alta',S_OEM:'Media',S_EPP:'Media',S_SEG:'Media',S_SAL:'Media',S_INST:'Media',S_PROV:'Baja',S_OUT:'Baja'}
ACC={S_MINE:'Identificar Gerente de Seguridad/SSO de la operación o proyecto; proponer piloto (rajo, planta, flota)',
 S_CONTR:'Identificar Jefe de Prevención/SSO; proponer piloto en frente de trabajo (perforación, transporte, campamento)',
 S_TECH:'Mapear oferta: evaluar integración (telemetría, comunicaciones, IA) o competencia',
 S_CONS:'Explorar alianza comercial / canal de implementación y formación en seguridad',
 S_OEM:'Explorar alianza con representante local / integración en equipos y talleres',
 S_EPP:'Evaluar como canal de distribución hacia mineras y contratistas',
 S_SEG:'Evaluar co-venta o integración de soluciones',
 S_SAL:'Evaluar alianza en salud ocupacional / fatiga',
 S_INST:'Networking institucional; identificar contacto y eventos',
 S_PROV:'Nutrir con contenido; sin acción comercial inmediata',
 S_OUT:'Sin acción comercial'}
CURK={key(n):s for s,ns in CUR.items() for n in ns}
NOTE_X={'TCV DESARROLLO DE SOLUCIONES':'Competidor directo: videovigilancia con IA contra fatiga, somnolencia y distracción en flotas.',
 'MAKERS SAS':'Software e IA a medida para minería (posible competidor/aliado).',
 'ASOCIACIÓN OBRERA MINERA ARGENTINA':'Sindicato minero (AOMA): actor clave en seguridad de trabajadores.',
 'CÁMARA MINERA DE SAN JUAN':'Gremio de las mineras de San Juan: puerta de entrada a operaciones.',
 'CASEMI':'Cámara de servicios mineros de San Juan (26 socios expositores).',
 'MINETECH':'Socio de Motorola Solutions y MSA en minería (comunicaciones y seguridad).',
 'SUMAR SALUD':'Salud ocupacional en sitios remotos (exámenes, campamentos, ambulancias).',
 'GLOBAL MINING - MAQUINARIAS RAWSON':'Brigadas de rescate para minería; ISO 45001.'}

rows=list(csv.DictReader(open(src,encoding='utf-8-sig')))
out_rows=[]; seen=set()
for x in rows:
    emp=clean(x['Empresa']); k=key(emp) or norm(emp)
    if k in seen: continue
    seen.add(k)
    seg=CURK.get(key(emp),S_PROV); pri=PRI[seg]
    ag=clean(x['Agrupacion']); agn=norm(ag)
    pais=PAIS_FIX.get(emp)
    if not pais:
        pais='Argentina'
        for d,p in DELEG.items():
            if d in agn: pais=p
    web=clean(x['SitioWeb']); email=clean(x['Email'])
    if 'panorama-minero' in web and 'PANORAMA' not in emp: web=''
    if 'panorama-minero' in email and 'PANORAMA' not in emp: email=''
    if 'linkedin' in web: web=' | '.join(w for w in web.split(' | ') if 'linkedin' not in w)
    stands=clean(x['Stands']); pab=clean(x['Pabellon']); ub=clean(x['Ubicacion'])
    loc='Exterior' if ub=='Exterior' else (f'Pabellón {pab}' if pab else '')
    desc=clean(x['Descripcion']) or clean(x['Productos'])
    notas='Expo San Juan Minera 2026.'
    if emp in NOTE_X: notas+=' '+NOTE_X[emp]
    if not email: notas+=' Sin email publicado.'
    cruce=[]
    for nm,m in CRMS:
        d=m.get(key(emp))
        if d: cruce.append(f"{nm} (ID {d['ID']})")
    if cruce: notas+=' También en: '+'; '.join(cruce)+'.'
    out_rows.append({'Nombre':'','Cargo':'','Empresa':emp,'País':pais,'Email':email,'LinkedIn':'','Origen':'Expo San Juan Minera 2026',
      'Estado':'Prospecto','Prioridad':pri,'Próxima acción':ACC[seg],'Notas':notas,'Segmento SafetyMind':seg,'Stand(s)':stands,'Ubicación':loc,
      'Agrupación':ag.split(': ',1)[-1],'Razón social':clean(x['RazonSocial']) if clean(x['RazonSocial']) not in ('Delegation','Camera','Group Industry Suppliers') else '',
      'Descripción':desc,'Web':web,'Cruce otros CRM':'; '.join(cruce) or 'No'})
porder={'Alta':0,'Media':1,'Baja':2}
out_rows.sort(key=lambda r:(porder[r['Prioridad']],ORDER.index(r['Segmento SafetyMind']),norm(r['Empresa'])))
COLS=['ID','Nombre','Cargo','Empresa','País','Email','LinkedIn','Origen','Estado','Prioridad','Próxima acción','Notas','Segmento SafetyMind','Stand(s)','Ubicación','Agrupación','Razón social','Descripción','Web','Cruce otros CRM']
wb=Workbook(); HF=PatternFill('solid',fgColor='1F3A5F'); HFont=Font(bold=True,color='FFFFFF')
def header(ws,cols,widths):
    ws.append(cols)
    for i,cn in enumerate(cols,1):
        c=ws.cell(1,i); c.fill=HF; c.font=HFont; c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1] if i-1<len(widths) else 16
    ws.freeze_panes='A2'
ws=wb.active; ws.title='Contactos'
header(ws,COLS,[6,20,20,40,14,32,14,22,12,10,50,55,48,22,12,26,30,60,32,26])
for i,r in enumerate(out_rows,1): ws.append([i]+[r[c] for c in COLS[1:]])
n=len(out_rows)+1
t=Table(displayName='Contactos',ref=f"A1:{get_column_letter(len(COLS))}{n}"); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
dv1=DataValidation(type='list',formula1='"Prospecto,Contactado,Reunión,Propuesta,Cliente,No interesado"',allow_blank=True); dv1.add(f'I2:I{n+500}')
dv2=DataValidation(type='list',formula1='"Alta,Media,Baja,Por calificar"',allow_blank=True); dv2.add(f'J2:J{n+500}')
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
ws2=wb.create_sheet('Oportunidades'); header(ws2,['ID oportunidad','ID contacto','Empresa','Necesidad','Solución propuesta','Etapa','Valor estimado','Probabilidad','Próxima acción','Fecha'],[14,12,40,40,40,14,14,12,40,12])
ws3=wb.create_sheet('Resumen'); header(ws3,['Evento/Fuente','País','Cantidad','',''],[55,22,12,12,12])
pc=collections.Counter(r['País'] for r in out_rows)
for i,(p,_) in enumerate(pc.most_common(),2): ws3.append(['Expo San Juan Minera 2026',p,f'=COUNTIFS(Contactos!$H:$H,$A{i},Contactos!$E:$E,$B{i})'])
last=ws3.max_row; ws3.append(['Total','',f'=SUM(C2:C{last})']); ws3.append([''])
def hdr(row,k):
    for cc in range(1,k+1): ws3.cell(row,cc).fill=HF; ws3.cell(row,cc).font=HFont
r0=ws3.max_row+1; ws3.append(['Segmento SafetyMind','Alta','Media','Baja','Total']); hdr(r0,5)
for s in ORDER:
    i=ws3.max_row+1; ws3.append([s]+[f'=COUNTIFS(Contactos!$M:$M,$A{i},Contactos!$J:$J,{col}${r0})' for col in 'BCD']+[f'=SUM(B{i}:D{i})'])
i=ws3.max_row+1; ws3.append(['Total']+[f'=SUM({cl}{r0+1}:{cl}{i-1})' for cl in 'BCDE']); ws3.append([''])
r1=ws3.max_row+1; ws3.append(['Estado (pipeline)','Cantidad']); hdr(r1,2)
for e_ in ['Prospecto','Contactado','Reunión','Propuesta','Cliente','No interesado']:
    i=ws3.max_row+1; ws3.append([e_,f'=COUNTIF(Contactos!$I:$I,A{i})'])
ws3.append(['']); r2=ws3.max_row+1; ws3.append(['Indicadores','Cantidad']); hdr(r2,2)
for lab,f in [('Registros originales',len(rows)),('Empresas / instituciones únicas','=COUNTA(Contactos!$D$2:$D$3000)'),('Con sede en Argentina','=COUNTIF(Contactos!$E$2:$E$3000,"Argentina")'),
  ('Con email','=COUNTIF(Contactos!$F$2:$F$3000,"?*")'),('Con web','=COUNTIF(Contactos!$S$2:$S$3000,"?*")'),('En otros CRM','=COUNTIF(Contactos!$T$2:$T$3000,"*(ID*")')]:
    ws3.append([lab,f])
ws4=wb.create_sheet('Eventos'); header(ws4,['Evento','País','Valor para SafetyMind'],[45,12,100])
for e in [['Directorio SSO 2026 – Cero Accidentes (22.ª ed.)','Perú','Mapa de ~670 proveedores SST en Perú (ver CRM_Directorio_SSO_2026).'],
 ['ExpoMina Perú 2026','Perú','Feria minera (ver CRM_ExpoMina_2026).'],['Exponor 2026 (Antofagasta)','Chile','Feria minera del norte de Chile (ver CRM_Exponor_2026).'],
 ['AquaSur 2026 (Puerto Montt)','Chile','Feria acuícola (ver CRM_AquaSur_2026).'],['FIDAE 2026 (Santiago)','Chile','Feria aeroespacial y de defensa (ver CRM_FIDAE_2026).'],
 ['Expo San Juan Minera 2026','Argentina','Feria minera de San Juan (cobre y oro): operaciones y proyectos (Veladero, Vicuña, Los Azules, El Pachón, Casposo), contratistas de perforación, voladura, transporte y campamentos, y cámaras (Cámara Minera, CASEMI, AOMA). Puerta de entrada a Argentina.']]: ws4.append(e)
wb.save(out)
print(len(rows),len(out_rows)); print(collections.Counter(r['Segmento SafetyMind'] for r in out_rows)); print(collections.Counter(r['Prioridad'] for r in out_rows))
print(collections.Counter(r['País'] for r in out_rows))
print('cruce',[(r['Empresa'],r['Cruce otros CRM']) for r in out_rows if r['Cruce otros CRM']!='No'])
print('not found',[x for ns in CUR.values() for x in ns if key(x) not in {key(r['Empresa']) for r in out_rows}])
json.dump(out_rows,open(out+'.json','w'),ensure_ascii=False)
