import csv, re, sys, json, collections, unicodedata
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
# usage: build_aquasur.py aquasur.csv CRM_Directorio_SSO_2026.xlsx CRM_ExpoMina_2026.xlsx CRM_Exponor_2026.xlsx out.xlsx
src, dirx, expox, exponx, out = sys.argv[1:6]
def clean(s): return re.sub(r'\s+',' ',(s or '').replace('\xa0',' ')).strip()
def norm(s): return unicodedata.normalize('NFKD',s.upper()).encode('ascii','ignore').decode()
def key(n):
    n=norm(n).lower()
    n=re.sub(r'\b(s\.?\s?a\.?\s?c\.?|s\.?\s?a\.?|e\.?i\.?r\.?l\.?|s\.?r\.?l\.?|ltda\.?|limitada|spa|sac|eirl|srl|peru|chile|del|de|y|cia|co|ltd|inc|corp|products|group|as)\b','',n)
    return re.sub(r'[^a-z0-9]','',n)
def load(path):
    ws=openpyxl.load_workbook(path)['Contactos']; h=[c.value for c in ws[1]]; m={}
    for r in ws.iter_rows(min_row=2,values_only=True):
        d=dict(zip(h,r))
        if d.get('Empresa'): m.setdefault(key(d['Empresa']),d)
    return m
dirmap=load(dirx); expomap=load(expox); exponmap=load(exponx)

RENAME={'VENDIDO':'COSESA','RESERVADO':'SUDVET','BRASIL':'Pabellón Brasil (MAPA)','NORWAY':'Innovation Norway',
 'ENDEAVOR':'Endeavor Aquatech Cluster','LUMIC (M-43)':'LUMIC','ECORILES (M-02)':'ECORILES'}
PAIS={'Corea , República de':'Corea del Sur','Netherlands':'Países Bajos','Croatia (Hrvatska)':'Croacia','New Zealand':'Nueva Zelanda'}
rows=list(csv.DictReader(open(src,encoding='utf-8-sig')))
comp=collections.OrderedDict()
for r in rows:
    raw=clean(r['Empresa'])
    name=RENAME.get(raw.upper(),RENAME.get(raw,raw))
    name=clean(re.sub(r'\s*\((?:[A-Z]-?\d+[A-Z]?|M-39E \(CREO\)|J-10|M-43|M-02|D-26|E-19|C-32|A-33)\)\s*$','',re.sub(r'\(M-39E \(CREO\)\)','',name)))
    k=key(name)
    c=comp.setdefault(k,{'Empresa':name,'stands':[],'giros':[],'desc':'','pais':PAIS.get(clean(r['Pais']),clean(r['Pais']))})
    st=clean(r['Stand'])
    for s in re.split(r'\s*,\s*',st):
        if s and s not in c['stands']: c['stands'].append(s)
    g=f"{clean(r['Giro1'])} › {clean(r['Giro2'])}"
    if g not in c['giros']: c['giros'].append(g)
    d=clean(r['Descripcion'])
    if d and d!='Sin descripción disponible.' and len(d)>len(c['desc']): c['desc']=d

S_MINE='Cliente final – productor acuícola'
S_CONTR='Cliente potencial – contratista / operación de alto riesgo'
S_TECH='Partner tecnológico / competidor'
S_CONS='Canal / aliado – consultoría SST'
S_OEM='Partner OEM – equipos acuícolas y marinos'
S_EPP='Canal / distribuidor EPP'
S_SEG='Aliado complementario – seguridad'
S_SAL='Aliado complementario – salud'
S_INST='Institución / gremio / pabellón país'
S_PROV='Proveedor industrial (equipos/insumos)'
# Lista curada a partir de nombre + descripción (el giro del CSV viene mal asignado en muchos casos)
CUR={
 S_MINE:['AQUACHILE','BLUMAR'],
 S_CONTR:['DIVING AUSTRAL','UNDERDEEP','Kafra Robotica spa','COMERCIAL INDUSTRIAL TERRAMAR SPA - WALBUSCH S.A','NACHIPA','FYT TRANSPORTES','Portuaria Cabo Froward','TOPTAINER','ASENAV','ASMAR SHIPYARDS','SOCIBER','BELATOR','TEKNICA','PELICANOS','TECH MARINE AND LOGISTICS SpA','SITRANS LTDA.','ALO RENTAL','SISDEF','NAVAL PRO SPA'],
 S_TECH:['AISBERG','ALFAPEOPLE','AQUABYTE','ACE AQUATEC','AKVA GROUP','BIOCEANOR','CAPTA HYDRO SPA','CONFIABILIDAD','DATAQU','ECTO, INC.','FRACTTAL','IMENCO','INCLUSIVE SECURITY','INNOVASEA','INNOVEX SPA','KNURO AS','LANIKS SPA.','LYTHIUM','MARINE CONNECTIONS','MARINE ROV','Marine View Robotics','Metaverso','METROCAPITAL','OXYGUARD COBALIA WATER','People Sync Spa','QYSEA Technology','ReelData AI','REY-AGUIRRE','ROBOTICA MARINA NEXSUB SPA','TecnoROV','Telemetria MS Spa','TIGABYTES','ULTRASEA','VORTEX SPA','ENTEL','GTD / TELSUR','INTERTELECOM','BlueDATAB','NAGU CORP','Wisely','ORCA TECNOLOGIA','TIDALX','GREENFOX MARINE','OPTOSCALE','PRAXIA SUR','NUCLEO ROJO','Quila TI SpA','BASELINE','Automation Anywhere','ELECTRONICA MARINA ITECH LIMITADA','IFM ELECTRONIC SPA','Ovun','DATASCOPE','FondeApp','ITG CHILE'],
 S_CONS:['APRO_X','Acuiestudios SpA','CINDER CAPACITACION','FERNANDEZ Y ALEUY CAPACITACION','OTIC CCC','SMART TRAINING','SINGULARES SPA','WSP GLOBAL','SGS CHILE','INTERTEK','LEYTON INSURANCE','UNIVERSIDAD DE CHILE','UNIVERSIDAD DE CONCEPCIÓN','INCAR','SAMS ENTERPRISE','STIM CHILE','Enroka','MANPOWER GROUP','PROFISUR'],
 S_OEM:['BADINOTTI','BOSSAQUA','CFLOW FISH','OCEIN AS','POSEIDON ACUICULTURA','ADRIA WINCH','JORLE INDUSTRIAL','FERNANDEZ JOVE','FRAMO','FINNING CHILE S.A','DISTRIBUIDORA CUMMINS CHILE S.A.','SALFA SUR','ABB S.A.','ROCKWELL AUTOMATION','ANCORA','MACKAY MARINE','INDUSMAR','PATAGONIA IMPORTACIONES','ELECTRORIDERS','FEEDING SYSTEMS SL','LUMIC','ROBERT BOSCH S.A.','BZ NAVAL ENGINEERING','FJORD MARITIME','NOGVA MOTOR','DERCOMAQ S.A.','HAOSAIL MACHINERY CHILE','PTG FRIONORDICA'],
 S_EPP:['APRO','AMSEC','FORUS S.A. - NORSEG','GSP CHILE','GSP-CHILE / DUNLOP PROTECTIVE FOOTWEAR','MACME','MARYUN SEGURIDAD INDUSTRIAL LIMITADA','MAX SERVICE','NOVA SEGURIDAD LIMITADA','SEGURYCEL S.A','SIRIAN','AGROFOR - BEKINA BOOTS','PONSA','Outfit_X'],
 S_SEG:['GANDARA CHILE','INGETECO','Gestion Control Seguridad S.A. - GCS','GSC CHILE - SEGURIDAD PRIVADA','RM SEGURIDAD CHILE','REPRESENTACIONES AEROTECH LIMITADA','LOCKERPLUS'],
 S_SAL:['MI DOCTOR'],
 S_INST:['ABRA - BRASIL','AMICHILE','AQUABENCH','BOMBEROS DE PUERTO VARAS','BOTE SALVAVIDAS','CAMARA CHILENO SUIZA DE COMERCIO A.G.','DANISH EMBASSY','Endeavor Aquatech Cluster','FACH - FUERZA AEREA DE CHILE','GLOBAL AQUA','HAVEXPO','MUNICIPALIDAD DE PUERTO MONTT','MUNICIPALIDAD DE PUERTO VARAS','PROCHILE','SERNAPESCA','DIRECTEMAR','CORFO - SALMON SUSTENTABLE','CLUB INNOVACION ACUICOLA','Innovation Norway','Pabellón Brasil (MAPA)','B2B MEDIA GROUP SPA','INFO SALMON','MUNDO ACUICOLA','RADIO PATAGONIA','REVISTA PUERTO A PUERTO','SALMONEXPERT','VC MAGAZINE','VERTICE TV','EL CONQUISTADOR','CERTIFIED HUMANE'],
}
CURK={key(n):s for s,ns in CUR.items() for n in ns}
TECHRX=r'inteligencia artificial|\bIA\b|\bAI\b|visi[oó]n (artificial|por computador)|computer vision|machine learning|\bIoT\b|\bROVs?\b|rob[oó]tica submarina|anal[ií]tica avanzada'
INSTRX=r'^(EMBAJADA|CAMARA|MUNICIPALIDAD|ASOCIACI)|REVISTA|MAGAZINE'
def seg_of(c):
    k=key(c['Empresa'])
    if k in CURK: return CURK[k]
    n=c['Empresa']; d=c['desc']; g=' '.join(c['giros'])
    if re.search(INSTRX,norm(n)): return S_INST
    if re.search(TECHRX,d) or re.search(TECHRX,n): return S_TECH
    if 'Medios especializados' in g: return S_INST
    if re.search(r'(?i)elementos de protecci[oó]n personal|\bEPP\b|calzado de seguridad|ropa de trabajo',d): return S_EPP
    if re.search(r'(?i)servicios? de buceo|wellboat|astillero',d) and 'Transporte acuícola' in g: return S_CONTR
    return S_PROV
ORDER=[S_MINE,S_CONTR,S_TECH,S_CONS,S_OEM,S_EPP,S_SEG,S_SAL,S_INST,S_PROV]
PRI={S_MINE:'Alta',S_CONTR:'Alta',S_TECH:'Alta',S_CONS:'Alta',S_OEM:'Media',S_EPP:'Media',S_SEG:'Media',S_SAL:'Media',S_INST:'Media',S_PROV:'Baja'}
ACC={S_MINE:'Mapear Gerente de Personas/SSOMA y centros de cultivo; proponer piloto (buceo, cosecha, wellboat)',
 S_CONTR:'Identificar Gerente HSE/Prevención en LinkedIn y proponer demo (faenas marítimas de alto riesgo)',
 S_TECH:'Mapear oferta: evaluar integración (cámaras, IoT, ROV) o competencia',
 S_CONS:'Explorar alianza comercial / reventa a sus clientes salmoneros',
 S_OEM:'Explorar integración con equipos (visión + telemetría en pontones, wellboats, jaulas)',
 S_EPP:'Evaluar como canal de distribución hacia acuicultura/industria',
 S_SEG:'Evaluar co-venta o integración de soluciones',
 S_SAL:'Evaluar alianza en salud ocupacional / fatiga',
 S_INST:'Networking institucional; identificar contacto',
 S_PROV:'Nutrir con contenido; sin acción comercial inmediata'}
out_rows=[]
for k,c in comp.items():
    seg=seg_of(c); pri=PRI[seg]
    d=dirmap.get(k); e=expomap.get(k); x=exponmap.get(k)
    src_=d or {}
    notas="AquaSur 2026."
    if len([s for s in c['stands'] if s!='Sin stand'])>1: notas+=" Varios stands (mayor inversión)."
    if c['stands']==['Sin stand']: notas+=" Sin stand asignado en el listado."
    cruce=[]
    if x: cruce.append(f"Exponor (ID {x['ID']})"); notas+=f" También expone en Exponor 2026 (ID {x['ID']})."
    if e: cruce.append(f"ExpoMina (ID {e['ID']})"); notas+=f" También expone en ExpoMina 2026 (ID {e['ID']})."
    if d: cruce.append(f"Directorio SSO (ID {d['ID']})"); notas+=f" También en CRM Directorio SSO 2026 (ID {d['ID']}); contacto heredado = filial Perú."
    if not d: notas+=" Sin contacto publicado."
    out_rows.append({'Nombre':src_.get('Nombre') or '','Cargo':src_.get('Cargo') or '','Empresa':c['Empresa'],'País':c['pais'],
      'Email':src_.get('Email') or '','LinkedIn':'','Origen':'AquaSur 2026','Estado':'Prospecto','Prioridad':pri,
      'Próxima acción':ACC[seg],'Notas':notas,'Segmento SafetyMind':seg,'Stand(s)':', '.join(c['stands']),
      'N° stands':len([s for s in c['stands'] if s!='Sin stand']),'Rubro AquaSur':' | '.join(c['giros']),'Descripción':c['desc'],
      'Teléfono(s)':src_.get('Teléfono(s)') or '','Cruce otros CRM':'; '.join(cruce) or 'No'})
porder={'Alta':0,'Media':1,'Baja':2}
out_rows.sort(key=lambda r:(porder[r['Prioridad']],ORDER.index(r['Segmento SafetyMind']),-r['N° stands'],norm(r['Empresa'])))
COLS=['ID','Nombre','Cargo','Empresa','País','Email','LinkedIn','Origen','Estado','Prioridad','Próxima acción','Notas','Segmento SafetyMind','Stand(s)','N° stands','Rubro AquaSur','Descripción','Teléfono(s)','Cruce otros CRM']
wb=Workbook(); HF=PatternFill('solid',fgColor='1F3A5F'); HFont=Font(bold=True,color='FFFFFF')
def header(ws,cols,widths):
    ws.append(cols)
    for i,cn in enumerate(cols,1):
        cell=ws.cell(1,i); cell.fill=HF; cell.font=HFont; cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.column_dimensions[get_column_letter(i)].width=widths[i-1] if i-1<len(widths) else 16
    ws.freeze_panes='A2'
ws=wb.active; ws.title='Contactos'
header(ws,COLS,[6,20,20,40,14,30,14,16,12,10,45,50,40,14,9,45,80,24,28])
for i,r in enumerate(out_rows,1):
    ws.append([i]+[r[c] for c in COLS[1:]]); ws.cell(ws.max_row,18).number_format='@'
n=len(out_rows)+1
t=Table(displayName='Contactos',ref=f"A1:{get_column_letter(len(COLS))}{n}"); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
dv1=DataValidation(type='list',formula1='"Prospecto,Contactado,Reunión,Propuesta,Cliente,No interesado"',allow_blank=True); dv1.add(f'I2:I{n+500}')
dv2=DataValidation(type='list',formula1='"Alta,Media,Baja,Por calificar"',allow_blank=True); dv2.add(f'J2:J{n+500}')
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
ws2=wb.create_sheet('Oportunidades'); header(ws2,['ID oportunidad','ID contacto','Empresa','Necesidad','Solución propuesta','Etapa','Valor estimado','Probabilidad','Próxima acción','Fecha'],[14,12,40,40,40,14,14,12,40,12])
ws3=wb.create_sheet('Resumen'); header(ws3,['Evento/Fuente','País','Cantidad','',''],[50,26,12,12,12])
pc=collections.Counter(r['País'] for r in out_rows)
for i,(p,_) in enumerate(pc.most_common(),2): ws3.append(['AquaSur 2026',p,f'=COUNTIFS(Contactos!$H:$H,$A{i},Contactos!$E:$E,$B{i})'])
last=ws3.max_row; ws3.append(['Total','',f'=SUM(C2:C{last})']); ws3.append([''])
def hdr(row,k):
    for cc in range(1,k+1): ws3.cell(row,cc).fill=HF; ws3.cell(row,cc).font=HFont
r0=ws3.max_row+1; ws3.append(['Segmento SafetyMind','Alta','Media','Baja','Total']); hdr(r0,5)
for s in ORDER:
    i=ws3.max_row+1; ws3.append([s]+[f'=COUNTIFS(Contactos!$M:$M,$A{i},Contactos!$J:$J,{col}${r0})' for col in 'BCD']+[f'=SUM(B{i}:D{i})'])
i=ws3.max_row+1; ws3.append(['Total']+[f'=SUM({cl}{r0+1}:{cl}{i-1})' for cl in 'BCDE']); ws3.append([''])
r1=ws3.max_row+1; ws3.append(['Estado (pipeline)','Cantidad']); hdr(r1,2)
for e_ in ['Prospecto','Contactado','Reunión','Propuesta','Cliente','No interesado']:
    i=ws3.max_row+1; ws3.append([e_,f'=COUNTIF(Contactos!$I:$I,A{i})'])
ws3.append(['']); r2=ws3.max_row+1; ws3.append(['Indicadores','Cantidad']); hdr(r2,2)
for lab,f in [('Registros originales',len(rows)),('Empresas únicas','=COUNTA(Contactos!$D$2:$D$3000)'),('También en Exponor 2026','=COUNTIF(Contactos!$S$2:$S$3000,"*Exponor*")'),('También en ExpoMina 2026','=COUNTIF(Contactos!$S$2:$S$3000,"*ExpoMina*")'),('También en Directorio SSO','=COUNTIF(Contactos!$S$2:$S$3000,"*Directorio*")'),('Con 2+ stands','=COUNTIF(Contactos!$O$2:$O$3000,">1")')]:
    ws3.append([lab,f])
ws4=wb.create_sheet('Eventos'); header(ws4,['Evento','País','Valor para SafetyMind'],[45,12,100])
ws4.append(['Directorio SSO 2026 – Cero Accidentes (22.ª ed.)','Perú','Mapa de ~670 proveedores SST en Perú (ver CRM_Directorio_SSO_2026).'])
ws4.append(['ExpoMina Perú 2026','Perú','Feria minera: mineras, contratistas, OEM, tecnología y proveedores (ver CRM_ExpoMina_2026).'])
ws4.append(['Exponor 2026 (Antofagasta)','Chile','Feria minera del norte de Chile (ver CRM_Exponor_2026).'])
ws4.append(['AquaSur 2026 (Puerto Montt)','Chile','Feria acuícola: salmoneras, contratistas marítimos (buceo, wellboats, astilleros), tecnología acuícola (IA, ROV, IoT) y proveedores; nueva vertical para SafetyMind.'])
wb.save(out)
print(len(rows),len(out_rows)); print(collections.Counter(r['Segmento SafetyMind'] for r in out_rows)); print(collections.Counter(r['Prioridad'] for r in out_rows)); print(pc.most_common(30))
print('cruce',collections.Counter(t.split(' (')[0] for r in out_rows for t in r['Cruce otros CRM'].split('; ') if t!='No'))
missing=[n for ns in CUR.values() for n in ns if key(n) not in comp and not n.endswith('_X')]
print('curated not found:',missing)
json.dump(out_rows,open(out+'.json','w'),ensure_ascii=False)
